"""Build an AUMOVIO deck from a JSON spec on the corporate template, using python-pptx (MIT).

Usage:
    python build_deck.py spec.json out.pptx

Text is measured with the real font files and shrunk to fit; bullet lists and
table rows that still do not fit are continued on extra slides.
Fit problems that cannot be solved are printed as WARN lines on stderr.
"""

import copy
import datetime
import json
import math
import os
import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "template.pptx"

ORANGE, ORANGE_TINT, PEACH = "FF4208", "FFF0EB", "FFB7A1"
PURPLE, PURPLE_LIGHT, BLUE, BLUE_LIGHT = "4827AF", "A08AE5", "18A9E7", "A3DDF5"
BLACK, WHITE = "111111", "FFFFFF"
CARD_STYLES = [(ORANGE, WHITE), (PURPLE, WHITE), (BLUE, WHITE), (PEACH, BLACK)]
CHART_COLORS = [ORANGE, PURPLE, BLUE, PEACH, PURPLE_LIGHT, BLUE_LIGHT]

# Content area of the template's white layouts, in inches.
LEFT, TOP, WIDTH, BOTTOM = 0.37, 1.47, 12.6, 6.75
HEIGHT = BOTTOM - TOP

LAYOUTS = {
    "title": "Cover (gradient)",
    "section": "Section (gradient) 01",
    "bullets": "Content with enumeration list",
    "content": "Title (white)",
    "closing": "End (gradient)",
}
FOOTER_TYPES = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)
SLIDE_TYPES = ("title", "section", "bullets", "cards", "table", "chart", "closing")

# Below this scale a list is continued on a new slide instead of shrunk further.
SPLIT_BELOW = 0.9
MIN_SCALE = 0.6

WARNINGS = []


def warn(msg):
    WARNINGS.append(msg)
    print(f"WARN {msg}", file=sys.stderr)


# ---------------------------------------------------------------- measuring

def _find_font(names):
    dirs = [
        Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "Windows" / "Fonts",
        Path("/Library/Fonts"),
        Path.home() / "Library" / "Fonts",
        Path("/usr/share/fonts/truetype/dejavu"),
    ]
    for d in dirs:
        for n in names:
            if (d / n).is_file():
                return str(d / n)
    return None


class Measure:
    """Text width and line height in points, from the brand font (Arial/DejaVu fallback)."""

    def __init__(self):
        self.fonts = {}
        for bold, names in (
            (False, ["AUMOVIOOffice-Regular.ttf", "arial.ttf", "Arial.ttf", "DejaVuSans.ttf"]),
            (True, ["AUMOVIOOffice-Bold.ttf", "arialbd.ttf", "Arial Bold.ttf", "DejaVuSans-Bold.ttf"]),
        ):
            path = _find_font(names)
            self.fonts[bold] = ImageFont.truetype(path, 1000) if path else None
        if self.fonts[False] is None:
            warn("no font file found for measuring; using a rough character-width estimate")

    def width(self, text, size, bold):
        font = self.fonts[bold] or self.fonts[False]
        if font is None:
            return len(text) * size * 0.55
        return font.getlength(text) * size / 1000

    def line_height(self, size, spacing, bold):
        font = self.fonts[bold] or self.fonts[False]
        em = 1.2 if font is None else max(sum(font.getmetrics()) / 1000, 1.15)
        return em * size * spacing


M = None


def para(text, size, bold=False, spacing=1.0, before=0, after=0, indent=0.0, level=None, color=None):
    return dict(text=str(text), size=size, bold=bold, spacing=spacing, before=before,
                after=after, indent=indent, level=level, color=color)


def count_lines(p, size, width_pt):
    space = M.width(" ", size, p["bold"])
    lines, cur = 0, None
    for chunk in p["text"].split("\n"):
        lines += 1
        cur = 0.0
        for word in chunk.split():
            w = M.width(word, size, p["bold"])
            if cur and cur + space + w <= width_pt:
                cur += space + w
                continue
            if cur:
                lines += 1
            # A word wider than the box wraps by character.
            extra = max(math.ceil(w / width_pt) - 1, 0)
            lines += extra
            cur = w - extra * width_pt
    return max(lines, 1)


def block_height(paras, scale, width_in):
    """Height in inches of paragraphs at the given font scale."""
    total = 0.0
    for i, p in enumerate(paras):
        size = p["size"] * scale
        width_pt = max((width_in - p["indent"]) * 72 * 0.97, 1)
        total += count_lines(p, size, width_pt) * M.line_height(size, p["spacing"], p["bold"])
        if i:
            total += p["before"] * scale
        if i < len(paras) - 1:
            total += p["after"] * scale
    return total / 72


def fit_scale(paras, width_in, height_in, min_scale=MIN_SCALE):
    """Largest scale (1.0 down to min_scale, 5% steps) at which paras fit; (scale, fits)."""
    scale = 1.0
    while scale >= min_scale - 1e-9:
        if block_height(paras, scale, width_in) <= height_in:
            return scale, True
        scale = round(scale - 0.05, 2)
    return min_scale, False


def paginate(groups, fits):
    """Greedily pack groups into pages so that fits(page) holds where possible."""
    pages, cur = [], []
    for g in groups:
        if cur and not fits(cur + [g]):
            pages.append(cur)
            cur = []
        cur.append(g)
    pages.append(cur)
    return pages


# ---------------------------------------------------------------- writing

def rgb(hex_str):
    return RGBColor.from_string(hex_str)


def write(tf, paras, scale):
    tf.clear()
    for i, p in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if p["level"] is not None:
            par.level = p["level"]
        par.alignment = PP_ALIGN.LEFT
        par.line_spacing = p["spacing"]
        par.space_before = Pt(p["before"] * scale) if i else Pt(0)
        par.space_after = Pt(p["after"] * scale)
        run = par.add_run()
        run.text = p["text"]
        run.font.size = Pt(round(p["size"] * scale, 1))
        run.font.bold = p["bold"]
        if p["color"]:
            run.font.color.rgb = rgb(p["color"])


def fill(shape, paras, width_in, height_in, where, min_scale=MIN_SCALE):
    scale, ok = fit_scale(paras, width_in, height_in, min_scale)
    if not ok:
        warn(f"{where}: text overflows even at {int(min_scale * 100)}% size - shorten it")
    write(shape.text_frame, paras, scale)
    return scale


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return box


def rect(slide, x, y, w, h, fill_hex, pad_x=0.25, pad_y=0.25):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill_hex)
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(pad_x)
    tf.margin_top = tf.margin_bottom = Inches(pad_y)
    return shp


def ph(slide, idx):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    return None


def drop(shape):
    if shape is not None:
        shape._element.getparent().remove(shape._element)


def dims(shape):
    return shape.width / 914400, shape.height / 914400


# ---------------------------------------------------------------- deck

class Deck:
    def __init__(self, spec):
        self.prs = Presentation(str(spec.get("template", TEMPLATE)))
        sld_ids = self.prs.slides._sldIdLst
        for sld_id in list(sld_ids):
            self.prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
        self.layouts = {l.name: l for l in self.prs.slide_layouts}
        self.footer = spec.get("meta", "")
        self.date = spec.get("date") or datetime.date.today().strftime("%B %Y")

    def new_slide(self, kind, notes=None):
        slide = self.prs.slides.add_slide(self.layouts[LAYOUTS[kind]])
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def where(self, what):
        return f"slide {len(self.prs.slides)} ({what})"

    def add_footer(self, slide):
        """add_slide() skips date/footer/number placeholders; copy them from the layout."""
        values = {PP_PLACEHOLDER.DATE: self.date, PP_PLACEHOLDER.FOOTER: self.footer}
        for lph in slide.slide_layout.placeholders:
            kind = lph.placeholder_format.type
            if kind not in FOOTER_TYPES or (kind in values and not values[kind]):
                continue
            el = copy.deepcopy(lph._element)
            el.nvSpPr.cNvPr.id = slide.shapes._next_shape_id
            slide.shapes._spTree.insert_element_before(el, "p:extLst")
            if kind in values:
                ph(slide, lph.placeholder_format.idx).text_frame.text = values[kind]

    def headline(self, slide, spec):
        """Orange category line (title placeholder) + black subheadline, one line each."""
        category, title = spec.get("category"), spec.get("title", "")
        top_text = category or title
        for idx, text in ((0, top_text), (14, title if category else None)):
            shape = ph(slide, idx)
            if text is None:
                drop(shape)
                continue
            w, h = dims(shape)
            fill(shape, [para(text, 30, spacing=0.8)], w, M.line_height(30, 0.8, False) / 72 + 0.02,
                 self.where("headline"))

    def content_slide(self, kind, spec, notes=True):
        slide = self.new_slide(kind, spec.get("notes") if notes else None)
        self.headline(slide, spec)
        self.add_footer(slide)
        return slide

    # ---- slide types

    def title(self, spec):
        slide = self.new_slide("title", spec.get("notes"))
        shape = ph(slide, 0)
        fill(shape, [para(spec["title"], 70, spacing=0.8)], *dims(shape), self.where("title"), 0.5)
        for idx, key, size in ((12, "subtitle", 30), (11, "org", 18), (13, "author", 18)):
            shape = ph(slide, idx)
            if spec.get(key):
                fill(shape, [para(spec[key], size)], *dims(shape), self.where(key))
            else:
                drop(shape)
        self.add_footer(slide)

    def section(self, spec):
        slide = self.new_slide("section", spec.get("notes"))
        shape = ph(slide, 0)
        fill(shape, [para(spec["title"], 70, spacing=0.8)], *dims(shape), self.where("section title"), 0.5)
        if spec.get("number"):
            ph(slide, 10).text_frame.text = str(spec["number"])
        else:
            drop(ph(slide, 10))
        if spec.get("subtitle"):
            box = textbox(slide, LEFT, 3.47, 8.6, 1.0)
            fill(box, [para(spec["subtitle"], 30, color=WHITE)], 8.6, 1.0, self.where("section subtitle"))

    def closing(self, spec):
        slide = self.new_slide("closing", spec.get("notes"))
        shape = ph(slide, 0)
        fill(shape, [para(spec.get("title", "Thank you"), 125, spacing=0.8)], *dims(shape),
             self.where("closing title"), 0.4)
        if spec.get("subtitle"):
            box = textbox(slide, LEFT, 3.6, WIDTH, 1.2)
            fill(box, [para(spec["subtitle"], 30, color=WHITE)], WIDTH, 1.2, self.where("closing subtitle"))

    def bullets(self, spec):
        groups = []
        for item in spec["bullets"]:
            if isinstance(item, str):
                item = {"text": item}
            if "heading" in item:
                groups.append([para(item["heading"], 18, True, 0.9, before=14, level=0)])
            elif "para" in item:
                groups.append([para(item["para"], 18, False, 0.9, before=8, level=1)])
            else:
                g = [para(item["text"], 18, False, 0.9, before=10, indent=0.295, level=2)]
                g += [para(s, 18, False, 0.9, before=4, indent=0.49, level=3) for s in item.get("sub", [])]
                groups.append(g)

        flat = [p for g in groups for p in g]
        scale, ok = fit_scale(flat, WIDTH, HEIGHT, SPLIT_BELOW)
        pages = [groups] if ok else paginate(
            groups, lambda page: block_height([p for g in page for p in g], 1.0, WIDTH) <= HEIGHT)

        for n, page in enumerate(pages):
            page_spec = dict(spec)
            if n:
                page_spec["title"] = f"{spec.get('title', '')} (cont.)"
            slide = self.content_slide("bullets", page_spec, notes=n == 0)
            fill(ph(slide, 16), [p for g in page for p in g], WIDTH, HEIGHT, self.where("bullets"))

    def cards(self, spec):
        slide = self.content_slide("content", spec)
        cards = spec["cards"]
        gap = 0.2
        w = (WIDTH - gap * (len(cards) - 1)) / len(cards)
        pad_x, pad_y = 0.25, 0.25
        inner_w = w - 2 * pad_x
        contents = [[para(c["heading"], 22, True, 0.9, after=10), para(c["body"], 18, False, 0.9)]
                    for c in cards]

        max_inner = HEIGHT - 0.2 - 2 * pad_y
        scale = min(fit_scale(c, inner_w, max_inner)[0] for c in contents)
        needed = max(block_height(c, scale, inner_w) for c in contents) + 2 * pad_y
        if needed > HEIGHT - 0.2 + 1e-6:
            warn(f"{self.where('cards')}: card text overflows even at {int(MIN_SCALE * 100)}% size")
        h = min(max(needed + 0.3, 2.2), HEIGHT - 0.2)
        for i, content in enumerate(contents):
            fill_hex, fg = CARD_STYLES[i % len(CARD_STYLES)]
            shp = rect(slide, LEFT + i * (w + gap), TOP + 0.2, w, h, fill_hex, pad_x, pad_y)
            for p in content:
                p["color"] = fg
            write(shp.text_frame, content, scale)

    def table(self, spec):
        cols, rows = spec["columns"], spec["rows"]
        weights = spec.get("col_widths") or [1] * len(cols)
        col_w = [WIDTH * wt / sum(weights) for wt in weights]
        pad_x, pad_y, min_h = 0.1, 0.06, 0.42

        def row_h(cells, scale, header=False):
            return max(min_h, max(block_height([para(c, 16, header)], scale, cw - 2 * pad_x)
                                  for c, cw in zip(cells, col_w)) + 2 * pad_y)

        def total(rows_, scale):
            return row_h(cols, scale, True) + sum(row_h(r, scale) for r in rows_)

        scale = 1.0
        while scale > SPLIT_BELOW and total(rows, scale) > HEIGHT:
            scale = round(scale - 0.05, 2)
        pages = [rows] if total(rows, scale) <= HEIGHT else paginate(
            [[r] for r in rows], lambda page: total([g[0] for g in page], 1.0) <= HEIGHT)
        if len(pages) > 1:
            scale = 1.0
            pages = [[g[0] for g in page] for page in pages]

        for n, page in enumerate(pages):
            page_spec = dict(spec)
            if n:
                page_spec["title"] = f"{spec.get('title', '')} (cont.)"
            slide = self.content_slide("content", page_spec, notes=n == 0)
            if total(page, scale) > HEIGHT:
                warn(f"{self.where('table')}: a single row is taller than the slide - shorten it")
            heights = [row_h(cols, scale, True)] + [row_h(r, scale) for r in page]
            shape = slide.shapes.add_table(len(page) + 1, len(cols), Inches(LEFT), Inches(TOP),
                                           Inches(WIDTH), Inches(sum(heights)))
            tbl = shape.table
            for c, cw in enumerate(col_w):
                tbl.columns[c].width = Inches(cw)
            for r, rh in enumerate(heights):
                tbl.rows[r].height = Inches(rh)
                header = r == 0
                for c in range(len(cols)):
                    cell = tbl.cell(r, c)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(ORANGE if header else (ORANGE_TINT if r % 2 == 0 else WHITE))
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.margin_left = cell.margin_right = Inches(pad_x)
                    cell.margin_top = cell.margin_bottom = Inches(pad_y)
                    text = cols[c] if header else page[r - 1][c]
                    write(cell.text_frame, [para(text, 16, header, color=WHITE if header else BLACK)], scale)

    def chart(self, spec):
        slide = self.content_slide("content", spec)
        ch = spec["chart"]
        kind = ch.get("type", "bar")
        data = CategoryChartData()
        data.categories = ch["categories"]
        for s in ch["series"]:
            data.add_series(s["name"], s["values"])
        xl = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
              "pie": XL_CHART_TYPE.PIE}[kind]
        width = 8.4 if spec.get("takeaway") else WIDTH
        chart = slide.shapes.add_chart(xl, Inches(LEFT), Inches(TOP), Inches(width), Inches(HEIGHT), data).chart
        chart.font.size = Pt(14)
        chart.has_title = False
        chart.has_legend = kind == "pie" or len(ch["series"]) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        plot.has_data_labels = True
        if kind == "pie":
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            for i, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(CHART_COLORS[i % len(CHART_COLORS)])
        else:
            for i, series in enumerate(plot.series):
                color = rgb(CHART_COLORS[i % len(CHART_COLORS)])
                if kind == "bar":
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = color
                else:
                    series.format.line.color.rgb = color
                    series.format.line.width = Pt(3)
                    series.marker.format.fill.solid()
                    series.marker.format.fill.fore_color.rgb = color
            chart.value_axis.major_gridlines.format.line.color.rgb = rgb("E5E5E5")
            chart.value_axis.format.line.fill.background()
            if kind == "bar":
                plot.gap_width = 60
                plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

        if spec.get("takeaway"):
            x = LEFT + width + 0.3
            w = LEFT + WIDTH - x
            panel = rect(slide, x, TOP, w, HEIGHT, ORANGE_TINT)
            content = [para(spec.get("takeaway_heading", "Key takeaway"), 20, True, 0.9, after=10, color=ORANGE),
                       para(spec["takeaway"], 18, False, 0.9, color=BLACK)]
            fill(panel, content, w - 0.5, HEIGHT - 0.5, self.where("takeaway"))

    def save(self, out_path):
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(out_path)


def build(spec, out_path):
    global M
    M = Measure()
    deck = Deck(spec)
    for s in spec["slides"]:
        if s["type"] not in SLIDE_TYPES:
            sys.exit(f"unknown slide type {s['type']!r}; expected one of {', '.join(SLIDE_TYPES)}")
        getattr(deck, s["type"])(s)
    deck.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    spec_file = Path(sys.argv[1])
    out = build(json.loads(spec_file.read_text(encoding="utf-8")), sys.argv[2])
    print(f"{out} ({len(WARNINGS)} warnings)")
