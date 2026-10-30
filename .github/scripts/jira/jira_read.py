"""
Jira Read - Unified ticket reader (API + Excel export support).

Provides unified interface to read Jira tickets from REST API or exported Excel files.
Automatically tries API first, then falls back to Excel export if needed.

Usage:
    python jira_read.py <TICKET_ID>
    python jira_read.py JZC-22658
    python jira_read.py JZC-22658 --fields summary,description,status
    python jira_read.py JZC-22658 --source api
    python jira_read.py JZC-22658 --source excel --excel-dir "E:\\JIRA"

Output:
    JSON object with ticket fields to stdout.

Configuration:
    Create .env file in the same directory with:
        JIRA_BASE_URL=https://ix.jira.automotive.cloud
        JIRA_EMAIL=your.email@company.com
        JIRA_API_TOKEN=your_api_token_here
    
    Or set these as environment variables.

Contact: Zhao Jianbo
Supported by: UX SW AI Community
"""

import sys
import os
import json
import argparse
import glob
from html.parser import HTMLParser

import jira_utils


# ---------------------------------------------------------------------------
# Excel file parser (HTML-based .xls format)
# ---------------------------------------------------------------------------
class _JiraHTMLParser(HTMLParser):
    """Parse HTML tables from Jira .xls exports."""
    def __init__(self):
        super().__init__()
        self.in_thead = False
        self.in_tbody = False
        self.in_td = False
        self.in_th = False
        self.headers = []
        self.rows = []
        self.current_row = []
        self.current_cell = ""

    def handle_starttag(self, tag, attrs):
        if tag == "thead":
            self.in_thead = True
        elif tag == "tbody":
            self.in_tbody = True
        elif tag == "th" and self.in_thead:
            self.in_th = True
        elif tag == "td" and self.in_tbody:
            self.in_td = True

    def handle_endtag(self, tag):
        if tag == "thead":
            self.in_thead = False
        elif tag == "tbody":
            self.in_tbody = False
        elif tag == "th" and self.in_thead:
            self.in_th = False
            self.headers.append(self.current_cell.strip())
            self.current_cell = ""
        elif tag == "td" and self.in_tbody:
            self.in_td = False
            self.current_row.append(self.current_cell.strip())
            self.current_cell = ""
        elif tag == "tr" and self.in_tbody and self.current_row:
            self.rows.append(self.current_row)
            self.current_row = []

    def handle_data(self, data):
        if self.in_th or self.in_td:
            self.current_cell += data


def _normalize_col(name):
    """Normalize column header to standard key."""
    ALIASES = {
        "issue key": "key", "key": "key", "issue_key": "key",
        "summary": "summary",
        "description": "description",
        "status": "status",
        "priority": "priority",
        "issue type": "issue_type", "issuetype": "issue_type", "type": "issue_type",
        "assignee": "assignee",
        "reporter": "reporter",
        "created": "created",
        "updated": "updated",
        "resolution": "resolution",
        "labels": "labels",
        "components": "components", "component/s": "components",
        "fix version/s": "fix_versions", "fix versions": "fix_versions",
        "affects version/s": "affects_versions", "affects versions": "affects_versions",
        "comment": "comments", "comments": "comments",
        "environment": "environment",
    }
    if name is None:
        return None
    name_lower = str(name).strip().lower()
    return ALIASES.get(name_lower, name_lower.replace(" ", "_").replace("/", "_"))


def _read_excel_xls(file_path, ticket_id):
    """Parse HTML-based .xls file and extract ticket record."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        html = f.read()

    parser = _JiraHTMLParser()
    try:
        parser.feed(html)
    except Exception as e:
        print(f"Warning: Error parsing {file_path}: {e}", file=sys.stderr)
        return None

    if not parser.headers or not parser.rows:
        return None

    normalized_headers = [_normalize_col(h) for h in parser.headers]

    for row in parser.rows:
        if len(row) != len(normalized_headers):
            continue
        record = dict(zip(normalized_headers, row))
        if record.get("key") == ticket_id:
            return record

    return None


def read_excel(ticket_id, excel_dir=None):
    """
    Find and read a Jira ticket from exported Excel file.
    Supports .xlsx (openpyxl) and .xls (HTML table) formats.
    """
    if excel_dir is None:
        excel_dir = os.getcwd()

    # Search for .xls files (HTML export)
    for xls_file in glob.glob(os.path.join(excel_dir, "*.xls")):
        record = _read_excel_xls(xls_file, ticket_id)
        if record:
            return record

    # Try .xlsx files (if openpyxl available)
    try:
        import openpyxl
        for xlsx_file in glob.glob(os.path.join(excel_dir, "*.xlsx")):
            try:
                wb = openpyxl.load_workbook(xlsx_file)
                ws = wb.active
                headers = [_normalize_col(cell.value) for cell in ws[1]]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if len(row) != len(headers):
                        continue
                    record = dict(zip(headers, row))
                    if record.get("key") == ticket_id:
                        return record
            except Exception as e:
                print(f"Warning: Error reading {xlsx_file}: {e}", file=sys.stderr)
                continue
    except ImportError:
        pass

    return None


# ---------------------------------------------------------------------------
# Read from API
# ---------------------------------------------------------------------------
def read_api(ticket_id, fields=None):
    """Fetch ticket from Jira REST API."""
    base_url, email, token, pat = jira_utils.get_config()
    extra_fields = fields.split(",") if fields else None
    return jira_utils.fetch_issue(base_url, email, token, pat, ticket_id, extra_fields)


# ---------------------------------------------------------------------------
# Unified interface
# ---------------------------------------------------------------------------
def read_ticket(ticket_id, source=None, fields=None, excel_dir=None):
    """
    Read a Jira ticket from API or Excel export.
    
    If source is None, tries API first, then falls back to Excel.
    
    Returns:
        dict: Ticket data
    """
    if source is None or source == "api":
        try:
            data = read_api(ticket_id, fields)
            return data
        except SystemExit:
            if source == "api":
                raise
            # Fall back to Excel
            print(f"API failed, trying Excel export...", file=sys.stderr)

    if source is None or source == "excel":
        data = read_excel(ticket_id, excel_dir)
        if data:
            return data
        print(f"Could not read {ticket_id} from Excel", file=sys.stderr)
        sys.exit(1)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(
        description="Read a Jira ticket from API or Excel export"
    )
    parser.add_argument("ticket_id", help="Jira ticket ID (e.g., JZC-22658)")
    parser.add_argument(
        "--fields", default=None,
        help="Comma-separated field names (API only; default: common fields)"
    )
    parser.add_argument(
        "--source", choices=["api", "excel", "auto"], default="auto",
        help="Read source: api, excel, or auto (try API first) [default: auto]"
    )
    parser.add_argument(
        "--excel-dir", default=None,
        help="Directory containing Excel exports [default: current directory]"
    )
    args = parser.parse_args()

    data = read_ticket(
        args.ticket_id,
        source=None if args.source == "auto" else args.source,
        fields=args.fields,
        excel_dir=args.excel_dir,
    )

    print(json.dumps(data, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
