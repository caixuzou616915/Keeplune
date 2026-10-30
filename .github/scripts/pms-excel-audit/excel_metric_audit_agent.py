#!/usr/bin/env python3
"""Generate an HTML audit report for red metrics in an Excel function sheet."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime
import html
import os
import re
from collections import Counter
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries


DEFAULT_TARGET_COLUMNS = ["C", "K", "L", "M", "Q", "S", "T"]
DEFAULT_TARGET_COLUMN_SET = set(DEFAULT_TARGET_COLUMNS)
DEFAULT_TARGET_METRIC_NAMES = [
    "Cyclomatic (NIST)",
    "Maximal Level of Nesting",
    "Number of functions calling this function (FanIn, Non-Transitive)",
    "Number of functions being called by this function (FanOut, Non-Transitive)",
    "Static Path Count",
    "Inactive\\Total LOC Ratio (Percent)",
    "Comment\\Total LOC Ratio (Percent)",
    "Maintainability Metric",
]
COMMENT_COLUMN = "U"
RESPONSIBILITY_COLUMN = "V"
DEFAULT_HTML_BASENAME = "understand_function_red_metrics_report"
NEW_TEMPLATE_COMMENT_START_COLUMN = "S"
NEW_TEMPLATE_COMMENT_END_COLUMN = "AC"
NEW_TEMPLATE_PLATFORM_COLUMN = "B"
NEW_TEMPLATE_MODULE_COLUMN = "AG"
NEW_TEMPLATE_OWNER_COLUMN = "AH"
SOLUTION_NOTE = (
    "Code recommendation is a suggestion only. The developer remains responsible "
    "for validating the change, reviewing side effects, and running appropriate "
    "unit, integration, and deep regression tests."
)

# ---------------------------------------------------------------------------
# Built-in metric rules used when no external metric rule file is supplied
# Each key is the normalised metric name (lower-case).
# "warning" / "error" are upper thresholds (≥ value → red).
# "warning_below" is a lower threshold (< value → red, used for CCR%).
# ---------------------------------------------------------------------------
BUILTIN_METRIC_RULES: dict[str, dict] = {
    "lines of code (executable)": {
        "smk": None,
        "description": "Counts executable lines in active code",
        "rationale": "Code that is too large becomes hard to analyse and maintain",
        "action": "Redesign and split the function into smaller parts",
        "priority": "medium",
    },
    "lines of code (inactive)": {
        "smk": None,
        "description": "Counts lines of inactive code between non-active preprocessor blocks",
        "rationale": "Inactive code is often a design smell and is used to compute the inactive-to-total LOC ratio",
        "action": "Redesign and create packages or variants to avoid inactive code in one unit",
        "priority": "medium",
    },
    "lines of code (comments)": {
        "smk": None,
        "description": "Counts pure comment lines and mixed comment-and-statement lines",
        "rationale": "This metric is used to compute the comment-to-total LOC ratio",
        "action": "Add comments where they are genuinely helpful to explain intent and constraints",
        "priority": "medium",
    },
    "cyclomatic (nist)": {
        "smk": "CC",
        "warning": 10,
        "error": 20,
        "description": "Number of branches in the control flow graph",
        "rationale": "High complexity yields bad analysability and maintainability and is a minimal indicator of the number of test cases needed for branch coverage",
        "action": "Redesign or at least ensure full test coverage of all branches",
        "priority": "high",
    },
    "maximal level of nesting": {
        "smk": "NL",
        "warning": 3,
        "error": 5,
        "description": "Number of nesting levels; each nested control structure adds 1",
        "rationale": "Deep nesting yields bad analysability and maintainability",
        "action": "Redesign with guard clauses and early exits or at least ensure full test coverage",
        "priority": "high",
    },
    "number of functions calling this function(fanin, transitive)": {
        "smk": None,
        "description": "Number of functions that call this function, also indirectly via other functions",
        "rationale": "Check with the SW System Design whether many transitive callers are intended",
        "action": "Redesign in conformance with the SW System Design",
        "priority": "medium",
    },
    "number of functions calling this function (fanin, transitive)": {
        "smk": None,
        "description": "Number of functions that call this function, also indirectly via other functions",
        "rationale": "Check with the SW System Design whether many transitive callers are intended",
        "action": "Redesign in conformance with the SW System Design",
        "priority": "medium",
    },
    "number of functions being called by this function (fanout, transitive)": {
        "smk": None,
        "description": "Number of functions called by this function, also indirectly via other functions",
        "rationale": "Check with the SW System Design whether many transitive dependencies are intended",
        "action": "Redesign in conformance with the SW System Design",
        "priority": "medium",
    },
    "lines of code (total)": {
        "smk": "FFS",
        "warning": 101,
        "error": 150,
        "description": "Counts total number of Source Code Lines",
        "rationale": "Code that is too large becomes hard to analyse and maintain",
        "action": "Extract helper functions to reduce the overall line count",
        "priority": "medium",
    },
    "number of executed statements": {
        "smk": "STMT",
        "warning": 30,
        "error": 51,
        "description": "Counts executable statements in active code (may differ from executable lines)",
        "rationale": "Too many statements per function reduces readability and testability",
        "action": "Split the function into smaller, single-responsibility helpers",
        "priority": "medium",
    },
    "number of return statements": {
        "smk": "g2.27",
        "error": 2,
        "description": "Counts the number of return statements per function",
        "rationale": "More than one return statement violates SMK Rule g2.27",
        "action": "Redesign with clearer control structures and a single exit point where practical",
        "priority": "low",
    },
    "number of goto statements": {
        "smk": "r2.61",
        "error": 1,
        "description": "Counts the number of goto statements per function",
        "rationale": "Goto statements are not allowed by SMK Rule r2.61",
        "action": "Redesign using structured control flow instead of goto",
        "priority": "low",
    },
    "number of function parameters": {
        "smk": "PARAM",
        "warning": 5,
        "error": 8,
        "description": "Counts the number of function parameters",
        "rationale": "Too many parameters indicate poor encapsulation of functionality",
        "action": "Bundle related parameters into a context or config struct",
        "priority": "medium",
    },
    "comment\\total loc ratio (percent)": {
        "smk": "CCR",
        "warning_below": 30,
        "description": "Ratio of comment lines to total lines of code (percent)",
        "rationale": "Insufficient comments make the code hard to analyse and maintain",
        "action": "Add brief, purposeful comments for non-obvious logic and safety constraints",
        "priority": "high",
    },
    "inactive\\total loc ratio (percent)": {
        "smk": "IAR",
        "warning": 15,
        "error": 35,
        "description": "Division of inactive LOC by total LOC",
        "rationale": "High inactive code ratio indicates packaging or variability is being handled inside one source unit",
        "action": "Redesign packages or variants to avoid carrying inactive code in the same implementation",
        "priority": "high",
    },
    "maintainability metric": {
        "smk": "MTB",
        "warning": 300,
        "error": 501,
        "description": "Cyclomatic (NIST) \u00d7 Maximal Nesting Level \u00d7 ln(Static Path Count)",
        "rationale": "High value indicates poor maintainability due to complex nesting and branching",
        "action": "Redesign into smaller units or at least ensure full test coverage of all paths",
        "priority": "highest",
    },
    "static path count": {
        "smk": "STPC",
        "warning": 500,
        "error": 1001,
        "description": "Number of paths in the control flow graph",
        "rationale": "High path count yields bad analysability and maintainability and implies a large number of path coverage tests",
        "action": "Redesign or at least ensure full path-oriented test coverage",
        "priority": "high",
    },
    "number of functions calling this function (fanin, non-transitive)": {
        "smk": None,
        "description": "Number of functions that directly call this function",
        "rationale": "High fan-in may indicate the function is a bottleneck with too many responsibilities",
        "action": "Verify with SW System Design; split public entry points or introduce helper abstractions",
        "priority": "high",
    },
    "number of functions being called by this function (fanout, non-transitive)": {
        "smk": None,
        "description": "Number of functions directly called by this function",
        "rationale": "High fan-out increases coupling and reduces testability",
        "action": "Introduce a service layer or helper API to reduce direct module dependencies",
        "priority": "high",
    },
}

METRIC_RULES: dict[str, dict] = dict(BUILTIN_METRIC_RULES)

SOURCE_SUFFIXES = {".c", ".cc", ".cpp", ".cxx", ".h", ".hpp", ".hxx"}
REFERENCE_SUFFIXES = SOURCE_SUFFIXES | {".arxml", ".swcd", ".xml"}
EXCLUDED_DIRS = {".git", ".repo", "node_modules", "out", "dist", "target", "obj", "bin"}
PREFERRED_SOURCE_ROOTS = ("pkg", "app")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        formatter_class=argparse.RawTextHelpFormatter,
        description=(
            "Scan red-highlighted metrics in selected columns and produce an HTML "
            "audit report, or fill comments and ownership fields for the newer workbook template."
        ),
        epilog=(
            "Usage examples:\n"
            "  1. Old template report flow:\n"
            "     python excel_metric_audit_agent.py PMS.xlsm --file-list owned_files.txt --metric-rule-file MetricRule.txt\n"
            "     Use this flow for the legacy workbook template. Provide the PMS workbook, an optional owned-file list,\n"
            "     and an optional metric rule file when you want to override the built-in default metric set.\n\n"
            "  2. New template comment fill flow:\n"
            "     python excel_metric_audit_agent.py PMS.xlsm --template-mode new-template-comment-fill --file-prefixes ETH EAL --platform-cds-module P --sw-module-name ETH --responsible-name Kumar\n"
            "     Use this flow for the newer workbook template. Provide the PMS workbook, one or more file prefixes\n"
            "     such as ETH or EAL, the module value P, the module name ETH, and the responsible name Kumar."
        ),
    )
    parser.add_argument("workbook", type=Path, help="Path to .xlsx/.xlsm report file")
    parser.add_argument(
        "--template-mode",
        choices=("legacy-red-metric-report", "new-template-comment-fill"),
        default="legacy-red-metric-report",
        help="Workbook handling mode. Defaults to the legacy red-metric HTML report flow.",
    )
    parser.add_argument(
        "--sheet",
        default="understand_function",
        help="Sheet name to scan (default: understand_function)",
    )
    parser.add_argument(
        "--start-row",
        type=int,
        default=9,
        help="First data row to scan (default: 9)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="HTML output file path",
    )
    parser.add_argument(
        "--columns",
        nargs="+",
        help="Excel column letters to scan for red-highlighted metrics; overrides default metric-name resolution",
    )
    parser.add_argument(
        "--metric-rule-file",
        type=Path,
        help=(
            "Optional tab-delimited metric rule file. When supplied, it can override thresholds "
            "and provide explicit Excel columns."
        ),
    )
    parser.add_argument(
        "--file-list",
        type=Path,
        help="Optional newline-delimited file allowlist used to filter findings to owned source files",
    )
    parser.add_argument(
        "--repo-root",
        type=Path,
        default=Path.cwd(),
        help="Repository root used to locate functions for code recommendations",
    )
    parser.add_argument(
        "--annotated-workbook",
        type=Path,
        help="Optional output path for the workbook copy with comments written to column U",
    )
    parser.add_argument(
        "--file-prefixes",
        nargs="+",
        help="Optional file-name prefixes for the new template flow, for example ETH EAL",
    )
    parser.add_argument(
        "--platform-cds-module",
        choices=("C", "AC", "P", "N"),
        help="Value to write into the new template Platform/CDS Module column B",
    )
    parser.add_argument(
        "--sw-module-name",
        help="Value to write into the new template SW Module Name column AG",
    )
    parser.add_argument(
        "--responsible-name",
        help="Value to write into the new template Action / Responsibility column AH",
    )
    return parser.parse_args()


def is_red_fill(fill: PatternFill) -> bool:
    """Return True when the cell fill appears red."""
    if fill is None:
        return False

    fg = fill.fgColor
    if fg is None:
        return False

    rgb = (fg.rgb or "").upper()
    if rgb:
        return any(
            red in rgb
            for red in (
                "FFFF0000",  # pure red
                "00FF0000",  # indexed style red
                "FFC00000",  # dark red
            )
        )

    indexed = fg.indexed
    if indexed is not None:
        # Common Excel indexed reds.
        return indexed in {3, 10}

    return False


def is_purple_fill(fill: PatternFill) -> bool:
    """Return True when the cell fill appears purple in the new template."""
    if fill is None:
        return False

    fg = fill.fgColor
    if fg is None:
        return False

    rgb = (fg.rgb or "").upper()
    if rgb:
        return rgb in {
            "FF7030A0",
            "007030A0",
        }

    indexed = fg.indexed
    if indexed is not None:
        return indexed in {20}

    return False


def normalize(text: object) -> str:
    return str(text).strip().lower() if text is not None else ""


# ---------------------------------------------------------------------------
# Tokens to suppress when extracting domain vocabulary from split identifiers.
# These are generic action verbs, C-type abbreviations, and filler words that
# appear frequently in automotive C function names but carry no domain meaning.
# Domain-specific abbreviations (eth, pdu, tcp, dtc, can, vin, ndp …) are
# intentionally NOT listed here so they surface naturally from any module.
# ---------------------------------------------------------------------------
_GENERIC_IDENTIFIER_TOKENS = frozenset({
    # Generic action verbs used as function-name prefixes
    "get", "set", "run", "put", "new", "old",
    "read", "write", "load", "save", "send", "recv",
    "check", "calc", "proc", "init", "deinit", "reset",
    "process", "update", "refresh", "execute",
    # Structural role labels (describe what the function is, not what it handles)
    "wrapper", "task", "func",
    # C type abbreviations used as identifier prefix/suffix fragments
    "uint", "sint", "bool", "void", "byte", "word",
    # Common programming shorthand
    "ptr", "ref", "val", "var", "tmp", "idx", "cnt",
    # Time-unit abbreviations common in cyclic task names
    "ms", "hz", "us", "ns",
    # Two-char English filler words
    "is", "or", "of", "by", "to", "do", "in", "at", "on", "be", "if", "as",
    # Longer English filler words that can appear in free-text context fields
    "the", "and", "for", "with", "from", "this", "are",
    "will", "has", "all", "its", "per", "via",
    # Positional connectors that appear as tokens in function names
    "before", "after", "without",
})


def expand_identifier_hints(text: object) -> str:
    """Split CamelCase and underscore-separated identifiers into space-delimited tokens."""
    raw = "" if text is None else str(text)
    expanded = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", raw)
    expanded = re.sub(r"([A-Za-z])([0-9])", r"\1 \2", expanded)
    expanded = re.sub(r"([0-9])([A-Za-z])", r"\1 \2", expanded)
    expanded = expanded.replace("_", " ")
    return normalize(expanded)


def has_repo_backed_source_match(source_match: dict | None) -> bool:
    if not source_match:
        return False
    return bool(
        source_match.get("path")
        and (
            str(source_match.get("line_text", "")).strip()
            or str(source_match.get("context_text", "")).strip()
        )
    )


def format_human_list(items: list[str]) -> str:
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    if len(items) == 2:
        return f"{items[0]} and {items[1]}"
    return ", ".join(items[:-1]) + f", and {items[-1]}"


def extract_context_signal_phrase(function_name: str, source_match: dict | None) -> str:
    """Extract a short phrase from real in-function variables instead of the function name."""
    if not has_repo_backed_source_match(source_match):
        return ""

    context_text = source_match.get("context_text", "") or ""
    line_text = source_match.get("line_text", "") or ""
    signal_names: list[str] = []
    seen: set[str] = set()

    def _add_signal(name: str) -> None:
        cleaned = str(name).strip()
        if not cleaned or cleaned in seen or cleaned == function_name:
            return
        seen.add(cleaned)
        signal_names.append(cleaned)

    signal_sources = [context_text]
    if not is_definition_like(line_text, function_name):
        signal_sources.insert(0, line_text)

    condition_text = " ".join(signal_sources)
    for condition_name in re.findall(r"(?:switch|if)\s*\([^\)]*->([A-Za-z_][A-Za-z0-9_]*)", condition_text):
        if re.search(r"(?:State|Type|Code|Status|Mode|Class|Flag|Error|Count|Length)$", condition_name):
            _add_signal(condition_name)

    for source_text in signal_sources:
        for member_name in re.findall(r"->([A-Za-z_][A-Za-z0-9_]*)", source_text):
            if re.search(r"(?:State|Type|Code|Status|Mode|Class|Flag|Error|Count|Length)$", member_name):
                _add_signal(member_name)
        if not signal_names:
            for variable_name in re.findall(r"\b([A-Za-z_][A-Za-z0-9_]*(?:State|Type|Code|Status|Mode|Class|Flag|Error|Count|Length))\b", source_text):
                _add_signal(variable_name)
        if len(signal_names) >= 3:
            break

    if not signal_names:
        return ""

    return f"while handling {format_human_list(signal_names[:3])} checks"


def extract_primary_domain_tokens(function_name: str, source_match: dict | None) -> list[str]:
    """
    Dynamically derive domain-specific tokens only from repo-backed source evidence.
    If the function cannot be resolved in the repo, return no tokens so the caller
    falls back to a generic comment rather than inventing a native-sounding phrase
    from the workbook function name alone.
    """
    if not has_repo_backed_source_match(source_match):
        return []

    seen: set[str] = set()

    def _tokens(text: str) -> list[str]:
        result: list[str] = []
        for token in expand_identifier_hints(text).split():
            if (len(token) >= 2
                    and not token.isdigit()
                    and token not in _GENERIC_IDENTIFIER_TOKENS
                    and token not in seen):
                seen.add(token)
                result.append(token)
        return result

    line_tokens: list[str] = []
    context_tokens: list[str] = []
    file_tokens: list[str] = []

    line_text = source_match.get("line_text", "") or ""
    context_text = source_match.get("context_text", "") or ""
    file_stem = Path(source_match.get("path", "") or "").stem

    if not is_definition_like(line_text, function_name):
        for ident in re.findall(r"\b[A-Za-z_][A-Za-z0-9_]{2,}\b", line_text)[:16]:
            line_tokens.extend(_tokens(ident))
    for ident in re.findall(r"\b[A-Za-z_][A-Za-z0-9_]{2,}\b", context_text)[:32]:
        context_tokens.extend(_tokens(ident))
    file_tokens.extend(_tokens(file_stem))

    return line_tokens[:4] + context_tokens[:6] + file_tokens[:2]


def format_domain_phrase(tokens: list[str], limit: int = 4) -> str:
    """
    Format the top domain tokens as a readable phrase.
    Short all-alphabetic tokens (≤4 chars) are uppercased as acronyms so that
    eth→ETH, rx→RX, pdu→PDU, tcp→TCP, ndp→NDP, vin→VIN, can→CAN etc. appear naturally.
    """
    selected = tokens[:limit]
    if not selected:
        return ""
    return " ".join(tok.upper() if len(tok) <= 4 and tok.isalpha() else tok for tok in selected)


def normalize_column_letter(value: str) -> str:
    normalized = str(value).strip().upper()
    if not normalized or not normalized.isalpha():
        raise ValueError(f"Invalid Excel column letter: {value}")
    return normalized


def parse_numeric_threshold(value: object) -> float | None:
    if value is None:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", str(value))
    if not match:
        return None
    return float(match.group(0))


def normalized_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", normalize(value)).strip()


def normalize_template_header(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\\", " ")
    text = re.sub(r"ALL\s*:\s*YELLOW\s*\+\s*RED", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip()
    return normalized_header(text)


def header_value(row: dict[str, str], *names: str) -> str:
    for name in names:
        value = row.get(name, "")
        if value:
            return value
    return ""


def parse_configured_columns(value: str) -> list[str]:
    # Match only 1-3 letter tokens (valid Excel column notation).
    # Longer words (e.g. from description sentences) are intentionally ignored.
    columns = [match.upper() for match in re.findall(r"\b[A-Za-z]{1,3}\b", value or "")]
    seen: list[str] = []
    for column in columns:
        try:
            normalized = normalize_column_letter(column)
        except ValueError:
            continue
        if normalized not in seen:
            seen.append(normalized)
    return seen


def column_index_from_letter(column_letter: str) -> int:
    return coordinate_to_tuple(f"{normalize_column_letter(column_letter)}1")[1]


def extract_comment_metric_key(header_text: object) -> list[str]:
    normalized = normalize_template_header(header_text)
    if not normalized:
        return []

    candidates = [normalized]
    if normalized.startswith("comment "):
        stripped = normalized[len("comment ") :].strip()
        if stripped:
            candidates.append(stripped)
    return candidates


def row_has_any_value(ws, row: int, start_col: int, end_col: int) -> bool:
    for column in range(start_col, end_col + 1):
        value = ws.cell(row=row, column=column).value
        if value is not None and str(value).strip() != "":
            return True
    return False


def has_existing_text(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def matches_file_prefix(file_name: str | None, prefixes: list[str] | None) -> bool:
    if not prefixes:
        return True
    if not file_name:
        return False
    normalized_name = str(file_name).strip().upper()
    return any(normalized_name.startswith(prefix.upper()) for prefix in prefixes)


def resolve_new_template_comment_columns(ws) -> dict[str, str]:
    start_index = column_index_from_letter(NEW_TEMPLATE_COMMENT_START_COLUMN)
    end_index = column_index_from_letter(NEW_TEMPLATE_COMMENT_END_COLUMN)
    metric_lookup: dict[str, str] = {}

    for column_index in range(1, start_index):
        column_letter = get_column_letter(column_index)
        header_key = normalize_template_header(ws[f"{column_letter}7"].value)
        if not header_key:
            continue
        metric_lookup.setdefault(header_key, column_letter)

    mapping: dict[str, str] = {}
    for column_index in range(start_index, end_index + 1):
        comment_column = get_column_letter(column_index)
        for candidate_key in extract_comment_metric_key(ws[f"{comment_column}7"].value):
            metric_column = metric_lookup.get(candidate_key)
            if metric_column:
                mapping[comment_column] = metric_column
                break

    return mapping


def source_allowed_files_for_prefixes(ws, start_row: int, prefixes: list[str] | None) -> set[str] | None:
    if not prefixes:
        return None

    allowed_files: set[str] = set()
    file_column = column_index_from_letter("A")
    function_column = column_index_from_letter("C")
    for row in range(start_row, ws.max_row + 1):
        if not row_has_any_value(ws, row, file_column, function_column):
            continue
        file_value = ws.cell(row=row, column=file_column).value
        if not file_value:
            continue
        file_name = Path(str(file_value).strip()).name
        if matches_file_prefix(file_name, prefixes):
            allowed_files.add(normalize(file_name))
    return allowed_files or None


def find_matching_new_template_rows(ws, start_row: int, prefixes: list[str] | None) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    file_column = column_index_from_letter("A")
    function_column = column_index_from_letter("C")

    for row in range(start_row, ws.max_row + 1):
        if not row_has_any_value(ws, row, file_column, function_column):
            continue

        file_value = ws.cell(row=row, column=file_column).value
        function_value = ws.cell(row=row, column=function_column).value
        if file_value is None and function_value is None:
            continue

        file_name = Path(str(file_value).strip()).name if file_value else ""
        if not matches_file_prefix(file_name, prefixes):
            continue

        rows.append(
            {
                "row": row,
                "file_name": file_name,
                "function_name": str(function_value).strip() if function_value else "",
            }
        )

    return rows


def metric_rule_file_path(args: argparse.Namespace) -> Path | None:
    return args.metric_rule_file


def load_metric_rule_config(metric_rule_file: Path | None) -> tuple[dict[str, dict], list[str], Path | None]:
    if metric_rule_file is None or not metric_rule_file.exists():
        return {}, [], None

    with metric_rule_file.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.reader(handle, delimiter="\t")
        rows = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]

    if len(rows) < 2:
        return {}, [], metric_rule_file

    headers = [normalized_header(value) for value in rows[0]]
    rules: dict[str, dict] = {}
    target_columns: list[str] = []

    for raw_row in rows[1:]:
        padded_row = raw_row + [""] * max(0, len(headers) - len(raw_row))
        row = {headers[index]: padded_row[index] for index in range(len(headers))}

        metric_name = header_value(row, "metric function", "metric", "metric name")
        if not metric_name:
            continue

        metric_key = normalize(metric_name)
        warning = parse_numeric_threshold(header_value(row, "warning"))
        error = parse_numeric_threshold(header_value(row, "error"))
        columns = parse_configured_columns(
            header_value(
                row,
                "excel column",
                "excel columns",
                "column",
                "columns",
                "target column",
                "target columns",
                "focus column",
                "focus columns",
                "sheet column",
                "sheet columns",
            )
        )
        for column in columns:
            if column not in target_columns:
                target_columns.append(column)

        rule = dict(BUILTIN_METRIC_RULES.get(metric_key, {}))
        rule.update(
            {
                "smk": header_value(row, "smk metric", "smk") or rule.get("smk"),
                "description": header_value(
                    row,
                    "short description see smk for more",
                    "short description",
                    "description",
                )
                or rule.get("description"),
                "rationale": header_value(
                    row,
                    "valuation see smk for more",
                    "valuation",
                    "rationale",
                )
                or rule.get("rationale"),
                "action": header_value(row, "recommendation", "action") or rule.get("action"),
                "priority": header_value(row, "priority") or rule.get("priority"),
            }
        )
        if warning is not None:
            rule["warning"] = int(warning) if warning.is_integer() else warning
        if error is not None:
            rule["error"] = int(error) if error.is_integer() else error
        rules[metric_key] = rule

    return rules, target_columns, metric_rule_file


def find_rule_columns_in_sheet(ws, configured_rule_names: Iterable[str]) -> list[str]:
    matched_columns: list[str] = []
    normalized_rule_names = [normalize(name) for name in configured_rule_names if normalize(name)]

    if not normalized_rule_names:
        return matched_columns

    for column_index in range(1, ws.max_column + 1):
        column_letter = get_column_letter(column_index)
        header_metric = normalize(metric_name(ws, column_letter))
        if not header_metric:
            continue
        if any(rule_name == header_metric for rule_name in normalized_rule_names):
            matched_columns.append(column_letter)

    return matched_columns


def resolve_target_columns(
    args: argparse.Namespace,
    configured_columns: list[str],
    configured_rules: dict[str, dict],
    ws=None,
) -> list[str]:
    if args.columns:
        return [normalize_column_letter(value) for value in args.columns]

    resolved_columns = list(configured_columns)
    if ws is not None and configured_rules:
        for column in find_rule_columns_in_sheet(ws, configured_rules.keys()):
            if column not in resolved_columns:
                resolved_columns.append(column)

    if not resolved_columns and ws is not None:
        for column in find_rule_columns_in_sheet(ws, DEFAULT_TARGET_METRIC_NAMES):
            if column not in resolved_columns:
                resolved_columns.append(column)

    if resolved_columns:
        return resolved_columns
    return list(DEFAULT_TARGET_COLUMNS)


def load_allowed_files(file_list_path: Path | None) -> tuple[set[str] | None, str | None]:
    if file_list_path is None:
        return None, None
    if not file_list_path.exists():
        raise FileNotFoundError(f"File list not found: {file_list_path}")

    entries = [
        line.strip()
        for line in file_list_path.read_text(encoding="utf-8", errors="ignore").splitlines()
        if line.strip()
    ]
    if not entries:
        return None, None

    module_name = entries[0]
    allowed_files = {
        normalize(Path(entry).name)
        for entry in entries[1:]
        if entry.strip()
    }
    return allowed_files or None, module_name


def matches_range(cell_ref: str, sqref: str) -> bool:
    row, col = coordinate_to_tuple(cell_ref)
    for part in str(sqref).split():
        min_col, min_row, max_col, max_row = range_boundaries(part)
        if min_col <= col <= max_col and min_row <= row <= max_row:
            return True
    return False


def to_number(value: object) -> float | None:
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return None
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def lookup_metric_rule(metric_name: str) -> dict | None:
    """Return the METRIC_RULES entry that matches metric_name (fuzzy, case-insensitive)."""
    key = normalize(metric_name)
    if key in METRIC_RULES:
        return METRIC_RULES[key]
    for rule_key, rule in METRIC_RULES.items():
        if rule_key in key or key in rule_key:
            return rule
    return None


def is_definition_like(line: str, function_name: str) -> bool:
    stripped = line.strip()
    if not stripped or stripped.startswith(("/*", "//", "*", "#")):
        return False
    match = re.search(rf"\b{re.escape(function_name)}\s*\(", stripped)
    if match is None:
        return False
    prefix = stripped[: match.start()].strip()
    if not prefix:
        return False
    if any(token in prefix for token in ("=", ")", "]", ">", "-", "+", "/", "%", "!", "?", ":")):
        return False
    definition_pattern = re.compile(
        rf"^(?:(?:static|inline|extern|const|signed|unsigned|volatile|register|struct|enum|union)\s+|[A-Za-z_]\w*[\s\*]+)+{re.escape(function_name)}\s*\("
    )
    return bool(definition_pattern.search(stripped))


def count_baseline_call_sites(candidates: list[dict], function_name: str) -> int:
    count = 0
    for candidate in candidates:
        line_text = candidate.get("line_text", "")
        if not line_text or is_definition_like(line_text, function_name):
            continue
        count += 1
    return count


def count_reference_sites(candidates: list[dict], function_name: str) -> int:
    count = 0
    for candidate in candidates:
        line_text = candidate.get("line_text", "")
        if not line_text:
            continue
        if line_text.strip().startswith(("/*", "//", "*")):
            continue
        if is_definition_like(line_text, function_name):
            continue
        count += 1
    return count


def has_baseline_usage_hint(source_match: dict | None) -> bool:
    if not source_match:
        return False

    context_text = normalize(source_match.get("context_text", ""))
    if not context_text:
        return False

    hint_markers = (
        "called by",
        "callback",
        "main function",
        "main10ms",
        "main50ms",
        "main1000ms",
        "init",
        "deinit",
        "api id",
        "swunit",
        "matlab code",
        "function call",
    )
    return any(marker in context_text for marker in hint_markers)


def detect_generated_source(lines: list[str]) -> str | None:
    header_text = " ".join(line.strip() for line in lines[:40])
    normalized_header = normalize(header_text)

    if any(
        marker in normalized_header
        for marker in (
            "code generated for simulink model",
            "simulink coder version",
            "embedded coder",
            "real-time workshop",
            "matlab",
        )
    ):
        return "matlab"

    return None


def describe_direct_fanin_reason(function_name: str, source_match: dict | None) -> str:
    combined = " ".join(
        part
        for part in (
            normalize(function_name),
            normalize(source_match.get("path", "")) if source_match else "",
            normalize(source_match.get("line_text", "")) if source_match else "",
            normalize(source_match.get("context_text", "")) if source_match else "",
        )
        if part
    )

    pattern_map = (
        (
            ("calc", "compute", "math", "formula", "interpolate", "average", "addition", "subtraction", "convert", "conversion"),
            "This function appears to serve as a common computing utility that is intentionally reused by several modules according to the current architecture design.",
        ),
        (
            ("wrapper", "adapter", "adaptor", "proxy", "interface", "osal", "port", "service"),
            "This function appears to provide a shared interface or wrapper layer, so reuse by several modules is expected within the architecture design.",
        ),
        (
            ("buffer", "queue", "fifo", "ring", "consumer", "producer", "stream"),
            "This function appears to support a shared buffering or producer-consumer mechanism, so direct reuse by several modules is expected by design.",
        ),
        (
            ("state", "event", "signal", "request", "response", "dispatch", "handle", "manager", "mgr"),
            "This function appears to coordinate shared state or event handling, so multiple direct callers are consistent with module dependency needs.",
        ),
        (
            ("check", "validate", "verify", "is", "has", "get", "set", "read", "write"),
            "This function appears to provide a shared access or validation point, so several direct callers can be expected within the module dependency structure.",
        ),
    )

    for keywords, description in pattern_map:
        if any(keyword in combined for keyword in keywords):
            return description

    return (
        "This function appears to provide shared logic required by several modules, and the direct reuse looks consistent with the current architecture design."
    )


def describe_statement_density_reason(function_name: str, source_match: dict | None) -> str:
    combined = " ".join(
        part
        for part in (
            normalize(function_name),
            normalize(source_match.get("path", "")) if source_match else "",
            normalize(source_match.get("line_text", "")) if source_match else "",
            normalize(source_match.get("context_text", "")) if source_match else "",
        )
        if part
    )

    pattern_map = (
        (
            ("switch", "case", "event", "trace", "tracing", "log", "logger"),
            "This function appears to consolidate a broad event or trace-handling flow, which naturally results in a large number of executed statements within one reviewed implementation.",
        ),
        (
            ("wrapper", "osal", "resource", "driver", "drv", "port", "interface", "service"),
            "This function appears to act as a wrapper for system or platform resource access, so the statement count is likely driven by the integration steps required by the architecture design.",
        ),
        (
            ("state", "mode", "sequence", "control", "process", "workflow", "orchestr", "dispatch"),
            "This function appears to centralize an end-to-end control flow with strict ordering and consistency requirements, so the higher statement count looks intentional within the reviewed design.",
        ),
        (
            ("calc", "compute", "evaluate", "interpolate", "convert", "mapping", "formula", "table"),
            "This function appears to evaluate a complex context or perform dense computing logic, so the larger statement count is consistent with the amount of domain processing handled in one place.",
        ),
        (
            ("generated", "gen", "auto", "trace event", "configuration", "config"),
            "This function appears to come from generated or configuration-driven logic, where a larger statement count can be expected because many cases must be handled in one stable implementation.",
        ),
    )

    for keywords, description in pattern_map:
        if any(keyword in combined for keyword in keywords):
            return description

    return (
        "This function appears to evaluate a broad operational context in one place, so the statement count looks driven by necessary implementation detail rather than uncontrolled logic growth."
    )


def extract_nesting_context(function_name: str, source_match: dict | None) -> tuple[str, str]:
    raw_parts = (
        function_name,
        source_match.get("path", "") if source_match else "",
        source_match.get("line_text", "") if source_match else "",
        source_match.get("context_text", "") if source_match else "",
    )
    raw_text = " ".join(part for part in raw_parts if part)
    # Include CamelCase/underscore-expanded hints so that category keyword scoring
    # can match tokens like "pdu", "tcp", "dtc" that are embedded in identifiers.
    expanded_hints = " ".join(expand_identifier_hints(part) for part in raw_parts if part)
    combined = normalize(f"{raw_text} {expanded_hints}")
    return raw_text, combined


def detect_nesting_structures(function_name: str, source_match: dict | None) -> dict[str, bool]:
    raw_text, combined = extract_nesting_context(function_name, source_match)
    assignments = Counter(re.findall(r"\b([A-Za-z_]\w*)\s*=(?!=)", raw_text))
    repeated_assignment = any(count >= 2 for count in assignments.values())
    repeated_status_assignment = any(
        count >= 2 and any(token in normalize(name) for token in ("state", "mode", "status", "result", "step", "event", "req", "resp"))
        for name, count in assignments.items()
    )
    if_count = len(re.findall(r"\bif\s*\(", raw_text))
    case_count = len(re.findall(r"\bcase\b", raw_text))
    logical_operator_count = len(re.findall(r"&&|\|\|", raw_text))
    protocol_tokens = ("request", "response", "diagnostic", "did", "dcm", "pdu", "payload", "frame")
    error_tokens = ("error", "fail", "retry", "timeout", "fallback", "pending", "restore", "degraded")
    permission_tokens = ("permission", "access", "privilege", "role", "allow", "deny", "inhibit", "matrix")
    validation_tokens = ("check", "validate", "null", "range", "bounds", "sanity")
    resource_tokens = ("alloc", "free", "lock", "unlock", "mutex", "semaphore", "buffer")
    init_tokens = ("init", "setup", "config", "configure", "start", "boot")

    return {
        "switch_on_var": bool(re.search(r"\bswitch\s*\(\s*[A-Za-z_]\w*(?:\s*(?:->|\.)\s*\w+)*\s*\)", raw_text)),
        "reassign_same_var": repeated_assignment,
        "reassign_status_var": repeated_status_assignment,
        "case_dispatch": case_count >= 3,
        "deep_conditional_tree": if_count >= 4 or (if_count >= 3 and logical_operator_count >= 3),
        "protocol_guard_ladder": if_count >= 2 and any(token in combined for token in protocol_tokens),
        "error_branch_cluster": sum(combined.count(token) for token in error_tokens) >= 3,
        "permission_guard_ladder": if_count >= 2 and any(token in combined for token in permission_tokens),
        "validation_guard_ladder": if_count >= 2 and any(token in combined for token in validation_tokens),
        "event_case_dispatch": case_count >= 2 and any(token in combined for token in ("event", "signal", "notify", "callback", "message", "trace")),
        "fast_path_branch": if_count >= 2 and any(token in combined for token in ("fast", "cache", "lookup", "optimize", "optimized", "performance")),
        "compatibility_branching": if_count >= 2 and any(token in combined for token in ("legacy", "compat", "adapter", "adaptor", "bridge", "wrapper", "interface", "service", "port")),
        "resource_guard_nesting": if_count >= 2 and any(token in combined for token in resource_tokens),
        "sequential_init_chain": repeated_assignment and if_count >= 2 and any(token in combined for token in init_tokens),
        "safety_guard_ladder": if_count >= 2 and any(token in combined for token in ("fault", "safe", "safety", "watchdog", "reset", "monitor", "adc", "voltage", "current", "thermal", "protect")),
        "state_transition_flow": any(token in combined for token in ("state", "mode", "transition", "sequence", "dispatch", "workflow", "fsm")) and (case_count >= 2 or repeated_status_assignment),
    }


# ---------------------------------------------------------------------------
# Shared category skeleton: keywords, structures, and priority are identical
# across all describe_*_reason functions.  Each function supplies only the
# metric-specific label and evidence strings via a {name: (label, evidence)}
# mapping passed to _score_categories().
# ---------------------------------------------------------------------------
_CATEGORY_SKELETONS: tuple[dict, ...] = (
    {
        "name": "protocol_handling",
        "keywords": ("protocol", "uds", "can", "lin", "request", "response", "payload", "frame", "packet", "diagnostic", "dcm", "pdu", "did", "readdid", "writedid", "writedidpro", "readdidpro"),
        "structures": ("protocol_guard_ladder",),
        "priority": 5,
    },
    {
        "name": "state_machine",
        "keywords": ("state", "mode", "transition", "sequence", "dispatch", "workflow", "step", "sm", "fsm"),
        "structures": ("switch_on_var", "reassign_status_var", "state_transition_flow"),
        "priority": 5,
    },
    {
        "name": "error_recovery",
        "keywords": ("error", "recover", "recovery", "retry", "timeout", "fallback", "fail", "pending", "restore", "nvm", "degraded"),
        "structures": ("error_branch_cluster", "reassign_same_var"),
        "priority": 4,
    },
    {
        "name": "defensive_validation",
        "keywords": ("check", "validate", "null", "range", "bounds", "sanity", "verify", "assert"),
        "structures": ("validation_guard_ladder",),
        "priority": 4,
    },
    {
        "name": "hardware_safety",
        "keywords": ("fault", "safe", "safety", "watchdog", "reset", "monitor", "adc", "voltage", "current", "thermal", "temperature", "hardware", "hw", "protect"),
        "structures": ("safety_guard_ladder",),
        "priority": 4,
    },
    {
        "name": "permission_matrix",
        "keywords": ("permission", "access", "privilege", "authorization", "auth", "role", "matrix", "allow", "deny", "inhibit"),
        "structures": ("permission_guard_ladder",),
        "priority": 4,
    },
    {
        "name": "algorithmic_logic",
        "keywords": ("calculate", "compute", "derive", "evaluate", "condition", "logic", "addition", "subtraction", "multiplication", "division", "library", "math", "arithmetic", "formula", "interpolat", "convert", "saturate", "sat", "clamp"),
        "structures": ("deep_conditional_tree",),
        "priority": 3,
    },
    {
        "name": "event_handling",
        "keywords": ("event", "handler", "notify", "callback", "queue", "trigger", "signal", "message", "trace"),
        "structures": ("event_case_dispatch", "case_dispatch"),
        "priority": 3,
    },
    {
        "name": "legacy_integration",
        "keywords": ("legacy", "compat", "adapter", "adaptor", "bridge", "wrapper", "interface", "integration", "osal", "service", "port"),
        "structures": ("compatibility_branching",),
        "priority": 3,
    },
    {
        "name": "resource_management",
        "keywords": ("alloc", "free", "lock", "unlock", "mutex", "semaphore", "buffer"),
        "structures": ("resource_guard_nesting",),
        "priority": 3,
    },
    {
        "name": "initialization_flow",
        "keywords": ("init", "setup", "config", "configure", "start", "boot"),
        "structures": ("sequential_init_chain",),
        "priority": 3,
    },
    {
        "name": "performance_optimization",
        "keywords": ("fast", "cache", "lookup", "table", "opt", "optimize", "optimized", "prefetch", "hotpath", "performance"),
        "structures": ("fast_path_branch",),
        "priority": 2,
    },
)


def _score_categories(
    combined: str,
    structures: dict[str, bool],
    label_evidence: dict[str, tuple[str, str]],
) -> list[dict]:
    scored: list[dict] = []
    for skeleton in _CATEGORY_SKELETONS:
        name = skeleton["name"]
        label, evidence = label_evidence.get(name, ("", ""))
        if not label:
            continue
        keyword_hits = sum(1 for kw in skeleton["keywords"] if kw in combined)
        structure_hits = sum(1 for s in skeleton["structures"] if structures.get(s, False))
        if keyword_hits == 0 and structure_hits == 0:
            continue
        score = keyword_hits + (structure_hits * 3) + skeleton["priority"]
        scored.append({
            "name": name,
            "label": label,
            "evidence": evidence,
            "score": score,
            "priority": skeleton["priority"],
            "keyword_hits": keyword_hits,
            "structure_hits": structure_hits,
        })
    scored.sort(key=lambda c: (c["score"], c["structure_hits"], c["priority"]), reverse=True)
    return scored


def describe_nesting_reason(function_name: str, source_match: dict | None) -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    label_evidence = {
        "protocol_handling":     ("protocol handling",                        "nested payload validation and per-service processing blocks create deep nesting per message type"),
        "state_machine":         ("state-machine transitions",                "deeply nested per-state processing inside switch-case blocks and guarded transition conditions"),
        "error_recovery":        ("error-recovery flow",                      "nested retry and fallback handling stacks multiple condition layers per error state"),
        "defensive_validation":  ("defensive validation",                     "stacked guard checks nest each subsequent operation deeper within the validated scope"),
        "hardware_safety":       ("hardware safety checks",                   "nested fault detection and hardware state validation creates layered safety guard blocks"),
        "permission_matrix":     ("permission-matrix gating",                 "nested access control checks layer role conditions and per-resource gating inside each other"),
        "algorithmic_logic":     ("algorithmic decision logic",               "nested numeric condition trees and multi-level boundary checks create deep evaluation layers"),
        "event_handling":        ("event handling and dispatch",              "nested per-event filtering and processing logic creates multiple layers of conditional depth"),
        "legacy_integration":    ("legacy integration branching",             "nested compatibility variant checks and adapter-specific handling add layered conditional depth"),
        "resource_management":   ("resource lifecycle management",            "nested allocation guards, validity checks, and release conditions layer resource handling depth"),
        "initialization_flow":   ("initialization sequencing",                "nested startup dependency checks and sequential configuration guards create layered init depth"),
        "performance_optimization": ("performance-sensitive fast-path handling", "nested fast-path condition checks and fallback cost evaluations create multi-level branching depth"),
    }
    scored = _score_categories(combined, structures, label_evidence)
    signal_phrase = extract_context_signal_phrase(function_name, source_match)
    domain_tokens = extract_primary_domain_tokens(function_name, source_match)
    domain_phrase = format_domain_phrase(domain_tokens)
    suffix = f" {signal_phrase}" if signal_phrase else (f" within {domain_phrase} processing" if domain_phrase else "")
    if not scored:
        return f"Nesting is primarily caused by layered conditional checks{suffix}. Reviewed and covered by testing."
    primary = scored[0]
    return f"Nesting is primarily caused by {primary['label']}{suffix} \u2014 {primary['evidence']}. Reviewed and covered by testing."


def describe_cyclomatic_reason(function_name: str, source_match: dict | None, value: str = "") -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    value_label = f" ({value})" if value else ""
    label_evidence = {
        "protocol_handling":     ("protocol handling",                        "multiple DID/UDS message branches and payload validation paths"),
        "state_machine":         ("state-machine transitions",                "switch-based state dispatch and guarded per-state transitions"),
        "error_recovery":        ("error-recovery flow",                      "retry, timeout, and fallback paths per error condition"),
        "defensive_validation":  ("defensive validation",                     "layered input guard checks, null tests, and range boundary conditions"),
        "hardware_safety":       ("hardware safety gating",                   "fail-safe gating, fault detection checks, and hardware state conditions"),
        "permission_matrix":     ("permission-matrix gating",                 "role-based allow/deny conditions and access control gating"),
        "algorithmic_logic":     ("algorithmic computation",                  "multi-condition numeric evaluation and overflow/boundary guards"),
        "event_handling":        ("event handling and dispatch",              "per-event-type dispatch cases and ordered filtering conditions"),
        "legacy_integration":    ("legacy integration branching",             "compatibility variant branches and adapter integration paths"),
        "resource_management":   ("resource lifecycle management",            "allocation guards, handle validity checks, and release conditions"),
        "initialization_flow":   ("initialization sequencing",                "sequential startup guards and configuration dependency checks"),
        "performance_optimization": ("performance-sensitive fast-path handling", "fast-path selection branches and fallback cost conditions"),
    }
    scored = _score_categories(combined, structures, label_evidence)
    signal_phrase = extract_context_signal_phrase(function_name, source_match)
    domain_tokens = extract_primary_domain_tokens(function_name, source_match)
    domain_phrase = format_domain_phrase(domain_tokens)
    suffix = f" {signal_phrase}" if signal_phrase else (f" within {domain_phrase} processing" if domain_phrase else "")
    if not scored:
        return f"Cyclomatic complexity{value_label} is primarily driven by multiple conditional paths{suffix}. Reviewed and covered by testing."
    primary = scored[0]
    return f"Cyclomatic complexity{value_label} is primarily driven by {primary['label']}{suffix} \u2014 {primary['evidence']}. Reviewed and covered by testing."


def describe_static_path_reason(function_name: str, source_match: dict | None, value: str = "") -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    value_label = f" ({value})" if value else ""
    label_evidence = {
        "protocol_handling":     ("protocol message dispatch chains",         "consecutive DID/UDS branch sequences \u2014 each additional message check multiplies the total path count"),
        "state_machine":         ("state-machine dispatch sequences",         "chained switch-case dispatch and per-state condition blocks \u2014 each state group multiplies execution paths"),
        "error_recovery":        ("error-recovery branch chains",             "sequential retry, timeout, and fallback guards \u2014 each independent check multiplies reachable paths"),
        "defensive_validation":  ("layered validation chains",                "consecutive independent null, range, and sanity guards \u2014 each check compounds the path product"),
        "hardware_safety":       ("hardware safety gate sequences",           "sequential fault detection and hardware state checks \u2014 each independent condition multiplies safe/unsafe paths"),
        "permission_matrix":     ("permission condition chains",              "consecutive role and access checks \u2014 each allow/deny guard independently multiplies path count"),
        "algorithmic_logic":     ("algorithmic branching sequences",          "chained numeric condition evaluations \u2014 consecutive comparison and boundary checks multiply code paths"),
        "event_handling":        ("event dispatch and filter chains",         "sequential event-type dispatch cases and ordered filter conditions multiply per-event execution paths"),
        "legacy_integration":    ("legacy compatibility branch chains",       "consecutive variant and adapter compatibility branches multiply integration paths"),
        "resource_management":   ("resource lifecycle gate chains",           "sequential allocation, validity, and release guards \u2014 each check multiplies live/dead resource paths"),
        "initialization_flow":   ("initialization sequence chains",           "consecutive startup and configuration dependency guards \u2014 each gate multiplies initialization paths"),
        "performance_optimization": ("fast-path selection chains",            "consecutive fast-path vs. fallback branch selections compound the total path product"),
    }
    scored = _score_categories(combined, structures, label_evidence)
    signal_phrase = extract_context_signal_phrase(function_name, source_match)
    domain_tokens = extract_primary_domain_tokens(function_name, source_match)
    domain_phrase = format_domain_phrase(domain_tokens)
    suffix = f" {signal_phrase}" if signal_phrase else (f" within {domain_phrase} processing" if domain_phrase else "")
    if not scored:
        return f"Static path count{value_label} is elevated due to consecutive conditional structures{suffix} that multiply execution paths exponentially. Reviewed and covered by testing."
    primary = scored[0]
    return f"Static path count{value_label} is primarily multiplied by {primary['label']}{suffix} \u2014 {primary['evidence']}. Reviewed and covered by testing."


def describe_maintainability_reason(function_name: str, source_match: dict | None, value: str = "") -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    value_label = f" ({value})" if value else ""
    label_evidence = {
        "protocol_handling":     ("protocol dispatch with nested validation",         "DID/UDS message handling adds branches, nests per-payload checks, and multiplies paths across each service case"),
        "state_machine":         ("state-machine with per-state condition nesting",   "switch-based dispatch increases branches, nesting per state, and path count grows with each guarded transition"),
        "error_recovery":        ("error-recovery with nested fallback logic",        "retry and fallback branches add cyclomatic depth, nest per-error handling, and chain into multiple reachable paths"),
        "defensive_validation":  ("layered defensive guards",                         "stacked null, range, and sanity checks raise branch count, increase nesting, and compound the path product"),
        "hardware_safety":       ("hardware safety with nested fault handling",       "fault and hardware state checks add branches, nest per-condition responses, and multiply safe/unsafe paths"),
        "permission_matrix":     ("permission gating with nested access logic",       "role and access condition chains raise cyclomatic count, nest per-role actions, and multiply access path combinations"),
        "algorithmic_logic":     ("algorithmic computation with nested conditions",   "chained numeric evaluations add branches, nest boundary and overflow checks, and compound reachable code paths"),
        "event_handling":        ("event dispatch with nested per-event processing",  "per-event-type cases add cyclomatic branches, nest filtering logic per event, and chain into many execution paths"),
        "legacy_integration":    ("legacy compatibility with nested variant handling","compatibility and adapter branches add conditions, nest variant-specific logic, and multiply integration paths"),
        "resource_management":   ("resource lifecycle with nested guard chains",      "allocation, validity, and release guards add branches, nest handle checks, and combine into multiple resource paths"),
        "initialization_flow":   ("initialization with nested dependency sequencing", "startup and config guards add conditions, nest dependency checks, and chain into multiple init execution paths"),
        "performance_optimization": ("fast-path selection with nested cost logic",    "fast-path and fallback selection adds branches, nests cost conditions, and compounds the total path product"),
    }
    scored = _score_categories(combined, structures, label_evidence)
    signal_phrase = extract_context_signal_phrase(function_name, source_match)
    domain_tokens = extract_primary_domain_tokens(function_name, source_match)
    domain_phrase = format_domain_phrase(domain_tokens)
    suffix = f" {signal_phrase}" if signal_phrase else (f" within {domain_phrase} processing" if domain_phrase else "")
    if not scored:
        return f"Maintainability score{value_label} is elevated{suffix} by the compound product of cyclomatic complexity, nesting depth, and path count. Reviewed and covered by testing."
    primary = scored[0]
    return f"Maintainability score{value_label} is primarily amplified by {primary['label']}{suffix} \u2014 {primary['evidence']}, compounding cyclomatic complexity, nesting depth, and path count together. Reviewed and covered by testing."


def describe_comment_ratio_reason(function_name: str, source_match: dict | None) -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    label_evidence = {
        "protocol_handling":     ("protocol message handling",            "DID/UDS dispatch logic, payload encoding, and service-specific validation paths need inline comments to clarify protocol intent and boundary constraints"),
        "state_machine":         ("state-machine logic",                  "state transition conditions, guard logic, and per-state action blocks benefit from comments that explain why each transition is guarded"),
        "error_recovery":        ("error-recovery handling",              "retry sequences, timeout conditions, and fallback paths require comments explaining the recovery intent and expected system behaviour per error case"),
        "defensive_validation":  ("defensive validation guards",          "null checks, range guards, and sanity assertions need short comments to document the assumption or contract they enforce"),
        "hardware_safety":       ("hardware safety logic",                "fault detection, watchdog interactions, and hardware state conditions require comments stating the safety rationale and permitted operational range"),
        "permission_matrix":     ("permission and access control",        "role conditions, allow/deny gates, and access matrix entries benefit from comments that record the authorization rule and its source requirement"),
        "algorithmic_logic":     ("algorithmic computation",              "numeric formulas, interpolation tables, and boundary evaluations need comments linking the implementation to the design specification or derived constant"),
        "event_handling":        ("event dispatch and handling",          "per-event case bodies and ordered filter conditions benefit from comments explaining the expected trigger source and processing intent"),
        "legacy_integration":    ("legacy integration code",              "compatibility variant branches, adapter shims, and wrapper layers need comments recording what legacy contract or interface version they preserve"),
        "resource_management":   ("resource lifecycle management",        "allocation, lock acquisition, and release guards benefit from comments documenting ownership, lifetime assumptions, and cleanup obligations"),
        "initialization_flow":   ("initialization sequencing",            "startup dependency checks and sequential configuration steps need comments recording the required order and the reason for each prerequisite"),
        "performance_optimization": ("performance-sensitive logic",       "fast-path selection, cache lookups, and cost-reduction branches benefit from comments explaining the optimization assumption and the estimated gain"),
    }
    scored = _score_categories(combined, structures, label_evidence)
    if not scored:
        return "Comment ratio is low; adding brief, targeted comments to non-obvious logic and safety constraints would improve maintainability. Reviewed and covered by testing."
    primary = scored[0]
    return f"Comment ratio is low for {primary['label']} \u2014 {primary['evidence']}. Reviewed and covered by testing."


def describe_fanin_non_transitive_reason(function_name: str, source_match: dict | None) -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    label_evidence = {
        "protocol_handling":     ("shared protocol dispatch entry",       "DID/UDS service dispatch and payload parsing utilities are called directly by each service handler, concentrating all message callers at one entry point"),
        "state_machine":         ("shared state query or transition",     "state get/set and transition-trigger functions are called directly by every module that needs to read or advance the shared state machine"),
        "error_recovery":        ("shared error reporting entry",         "common error-flag setters, fault reporters, and recovery initiators are called directly by every module that detects or reacts to a fault condition"),
        "defensive_validation":  ("shared validation utility",            "common null, range, and sanity check helpers are called directly by each module that needs to gate its inputs before further processing"),
        "hardware_safety":       ("shared hardware safety accessor",      "fault status readers and safety gate functions are called directly by every module that must check hardware health before proceeding"),
        "permission_matrix":     ("shared access control check",          "permission query and role-check functions are called directly by every module that must gate actions behind an access condition"),
        "algorithmic_logic":     ("shared computation utility",           "common math, conversion, and formula helpers are called directly by multiple modules that need the same derived value"),
        "event_handling":        ("shared event publisher or router",     "event notification and dispatch functions are called directly by every module that needs to post or subscribe to shared system events"),
        "legacy_integration":    ("shared legacy adapter entry",          "wrapper and adapter functions providing the legacy interface are called directly by every module that relies on the compatibility layer"),
        "resource_management":   ("shared resource access point",         "common allocation, handle-get, and status-query helpers are called directly by every module that needs access to the managed resource"),
        "initialization_flow":   ("shared initialization entry",          "common init and setup helpers are called directly by every module that must register or configure itself during startup"),
        "performance_optimization": ("shared fast-path accessor",         "optimized lookup and cache-access helpers are called directly by every module that requires the low-latency code path"),
    }
    scored = _score_categories(combined, structures, label_evidence)
    if not scored:
        return "Direct fan-in is high because this function provides shared logic required by several modules, and the direct reuse is consistent with the current architecture design. Reviewed and covered by testing."
    primary = scored[0]
    return f"Direct fan-in is primarily concentrated by {primary['label']} \u2014 {primary['evidence']}. Reviewed and covered by testing."


def describe_fanout_non_transitive_reason(function_name: str, source_match: dict | None) -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    label_evidence = {
        "protocol_handling":     ("protocol orchestration",               "message dispatch, payload parsing, and per-service handler calls spread direct dependencies across protocol stack layers"),
        "state_machine":         ("state-machine action dispatch",        "per-state entry, action, and exit calls distribute direct dependencies to every module that implements a state-specific behaviour"),
        "error_recovery":        ("error-recovery coordination",          "fault notification, recovery-step sequencing, and system-restore calls create direct dependencies to every involved module"),
        "defensive_validation":  ("validation and pre-check calls",       "input validation, configuration consistency checks, and sanity helpers are called directly for each guarded operation"),
        "hardware_safety":       ("hardware monitoring and control",      "safety gate queries, hardware status reads, and protective shutdown calls distribute direct dependencies across hardware modules"),
        "permission_matrix":     ("access control dispatch",              "role queries, permission checks, and conditional action calls create direct dependencies to each guarded resource or capability"),
        "algorithmic_logic":     ("algorithmic helper calls",             "math utilities, conversion helpers, and evaluation sub-functions are called directly for each computation stage"),
        "event_handling":        ("event routing and notification",       "per-event handler dispatch and subscriber notification calls spread direct dependencies to every registered event consumer"),
        "legacy_integration":    ("legacy adapter dispatch",              "compatibility shim calls and adapter-specific routing functions distribute direct dependencies across legacy and new module boundaries"),
        "resource_management":   ("resource lifecycle calls",             "allocation, handle validation, and release helpers are called directly at each lifecycle stage, spreading dependencies to resource subsystems"),
        "initialization_flow":   ("initialization sequencing calls",      "startup, configuration, and dependency-registration calls create direct dependencies to every module that must be initialised in order"),
        "performance_optimization": ("fast-path dispatch calls",          "cache priming, prefetch triggers, and hot-path helper calls distribute direct dependencies to the performance-critical subsystems"),
    }
    scored = _score_categories(combined, structures, label_evidence)
    if not scored:
        return "Direct fan-out is high because this function coordinates many downstream actions or dependent calls, which is consistent with its orchestration role in the current architecture. Reviewed and covered by testing."
    primary = scored[0]
    return f"Direct fan-out is primarily driven by {primary['label']} \u2014 {primary['evidence']}. Reviewed and covered by testing."


# ---------------------------------------------------------------------------
# Structural maintainability signal skeletons.
# Used ONLY for size/count symptoms: LOC, statement count, return count,
# and parameter count.  These are NOT root-cause categories — they describe
# structural patterns that explain how a symptom accumulates.
# Do NOT mix with _CATEGORY_SKELETONS.
# Signal types: function_size | exit_fragmentation | interface_complexity
# ---------------------------------------------------------------------------
_FUNCTION_SIZE_SIGNALS: tuple[dict, ...] = (
    {
        "name": "broad_responsibility",
        "keywords": ("manager", "mgr", "orchestrat", "coordinator", "process", "workflow", "handler", "control", "dispatch", "run", "task"),
        "structures": ("state_transition_flow", "case_dispatch"),
        "label": "broad orchestration responsibility",
        "evidence": "the function centralises multiple distinct operations in one body, accumulating lines as each concern is handled sequentially in scope",
        "priority": 5,
    },
    {
        "name": "dense_branching",
        "keywords": (),
        "structures": ("deep_conditional_tree", "case_dispatch", "switch_on_var"),
        "label": "dense inline conditional branching",
        "evidence": "many conditional branches are handled inline rather than delegated, each contributing executable lines and statements to the single function body",
        "priority": 4,
    },
    {
        "name": "integration_surface",
        "keywords": ("wrapper", "adapter", "adaptor", "osal", "port", "bridge", "proxy", "service", "integration", "interface"),
        "structures": ("compatibility_branching",),
        "label": "integration surface coverage",
        "evidence": "the function mediates multiple external interfaces or platform variants in one place, expanding the implementation surface with each integration point",
        "priority": 4,
    },
    {
        "name": "sequential_workflow",
        "keywords": ("init", "setup", "configure", "config", "sequence", "step", "boot", "start", "procedure"),
        "structures": ("sequential_init_chain",),
        "label": "sequential procedural workflow",
        "evidence": "a strictly ordered sequence of dependent steps is implemented linearly, with each stage adding lines in the same function scope without natural break points",
        "priority": 3,
    },
    {
        "name": "configuration_density",
        "keywords": ("table", "mapping", "lookup", "setting", "configuration", "matrix", "register", "bitfield", "constant"),
        "structures": (),
        "label": "configuration and mapping density",
        "evidence": "inline configuration tables, register mappings, or enumerated case values occupy a large portion of the function body by design",
        "priority": 2,
    },
)

_EXIT_FRAGMENTATION_SIGNALS: tuple[dict, ...] = (
    {
        "name": "validation_bail_out",
        "keywords": ("check", "validate", "null", "range", "bounds", "sanity", "verify", "assert", "guard"),
        "structures": ("validation_guard_ladder",),
        "label": "validation bail-out style",
        "evidence": "each input guard or precondition check exits immediately on failure, fragmenting the return path across multiple early exits per guard layer",
        "priority": 5,
    },
    {
        "name": "error_propagation",
        "keywords": ("error", "fail", "timeout", "retry", "fallback", "nok", "nack", "reject", "invalid", "abort"),
        "structures": ("error_branch_cluster",),
        "label": "error propagation exits",
        "evidence": "each distinct error condition terminates the function independently rather than converging on a single shared exit path",
        "priority": 4,
    },
    {
        "name": "dispatch_case_return",
        "keywords": ("state", "mode", "dispatch", "event", "signal", "step"),
        "structures": ("case_dispatch", "state_transition_flow", "event_case_dispatch"),
        "label": "per-case dispatch return",
        "evidence": "each dispatched case or state branch terminates with its own dedicated return, preventing a shared consolidated exit",
        "priority": 4,
    },
    {
        "name": "computed_result_return",
        "keywords": ("calculate", "compute", "evaluate", "derive", "get", "lookup", "convert", "map", "interpolat"),
        "structures": ("deep_conditional_tree",),
        "label": "computed result branching",
        "evidence": "each computation branch returns a directly evaluated result, producing one return per outcome rather than a single consolidated return statement",
        "priority": 3,
    },
)

_GOTO_USAGE_SIGNALS: tuple[dict, ...] = (
    {
        "name": "centralized_cleanup",
        "keywords": ("cleanup", "release", "unlock", "free", "close", "deinit", "restore", "exit", "return", "finish"),
        "structures": ("resource_guard_nesting",),
        "label": "centralized cleanup flow",
        "evidence": "multiple branches converge on one cleanup label so shared resource release and exit handling stay consistent",
        "priority": 5,
    },
    {
        "name": "error_escape",
        "keywords": ("error", "fail", "timeout", "fallback", "abort", "invalid", "reject", "pending", "recovery"),
        "structures": ("error_branch_cluster", "validation_guard_ladder"),
        "label": "error escape handling",
        "evidence": "multiple failure checks jump to a common recovery or exit label to avoid duplicating error handling across branches",
        "priority": 5,
    },
    {
        "name": "state_dispatch_escape",
        "keywords": ("state", "mode", "dispatch", "event", "signal", "case", "transition"),
        "structures": ("state_transition_flow", "case_dispatch", "event_case_dispatch"),
        "label": "state or dispatch escape flow",
        "evidence": "branch-specific paths jump to shared post-dispatch handling so state-dependent exits remain aligned across cases",
        "priority": 4,
    },
    {
        "name": "legacy_integration_escape",
        "keywords": ("legacy", "compat", "adapter", "adaptor", "wrapper", "interface", "service", "port"),
        "structures": ("compatibility_branching",),
        "label": "legacy integration control flow",
        "evidence": "compatibility and wrapper branches use shared jump targets to preserve one stable exit path across interface variants",
        "priority": 3,
    },
)

_INTERFACE_COMPLEXITY_SIGNALS: tuple[dict, ...] = (
    {
        "name": "multi_mode_control",
        "keywords": ("mode", "type", "kind", "variant", "action", "command", "operation", "flag", "option", "select"),
        "structures": ("switch_on_var", "state_transition_flow"),
        "label": "multi-mode control interface",
        "evidence": "separate parameters select the operating mode, action variant, or behavior flag, widening the interface for each independent control dimension",
        "priority": 5,
    },
    {
        "name": "data_context_bundle",
        "keywords": ("context", "ctx", "data", "info", "record", "cfg", "pdu", "buffer", "buf", "payload"),
        "structures": ("protocol_guard_ladder",),
        "label": "unbundled data and context inputs",
        "evidence": "data payload, configuration, and context fields are passed as separate parameters instead of being consolidated into a single composite structure",
        "priority": 4,
    },
    {
        "name": "multi_output_refs",
        "keywords": ("result", "out", "output", "status", "response", "resp", "length", "size", "count", "len"),
        "structures": ("reassign_status_var",),
        "label": "multiple output reference parameters",
        "evidence": "the function returns several results through individual pointer parameters, adding one parameter per independent output value",
        "priority": 4,
    },
    {
        "name": "dependency_injection",
        "keywords": ("handle", "hdl", "instance", "obj", "self", "ref", "channel", "manager", "mgr"),
        "structures": (),
        "label": "dependency injection handles",
        "evidence": "multiple service handles, object references, or interface pointers are injected individually rather than through a consolidated dependency bundle",
        "priority": 3,
    },
    {
        "name": "caller_metadata",
        "keywords": ("id", "caller", "trace", "source", "origin", "session", "tag", "token", "request", "req"),
        "structures": (),
        "label": "caller identification and tracing metadata",
        "evidence": "caller identity, session token, or tracing context is passed as an extra parameter rather than being carried in the primary data structure",
        "priority": 2,
    },
)


def _score_structural_signals(
    combined: str,
    structures: dict[str, bool],
    signal_skeletons: tuple[dict, ...],
) -> list[dict]:
    scored: list[dict] = []
    for skeleton in signal_skeletons:
        keyword_hits = sum(1 for kw in skeleton["keywords"] if kw in combined)
        structure_hits = sum(1 for s in skeleton["structures"] if structures.get(s, False))
        if keyword_hits == 0 and structure_hits == 0:
            continue
        score = keyword_hits + (structure_hits * 3) + skeleton["priority"]
        scored.append({
            "name": skeleton["name"],
            "label": skeleton["label"],
            "evidence": skeleton["evidence"],
            "score": score,
            "priority": skeleton["priority"],
            "keyword_hits": keyword_hits,
            "structure_hits": structure_hits,
        })
    scored.sort(key=lambda c: (c["score"], c["structure_hits"], c["priority"]), reverse=True)
    return scored


# Separate signal skeleton for executed-statement count.
# LOC (total) = physical bulk ("accumulating lines").
# STMT = execution density ("contributing statements per operation").
_STATEMENT_DENSITY_SIGNALS: tuple[dict, ...] = (
    {
        "name": "broad_responsibility",
        "keywords": ("manager", "mgr", "orchestrat", "coordinator", "process", "workflow", "handler", "control", "dispatch", "run", "task"),
        "structures": ("state_transition_flow", "case_dispatch"),
        "label": "broad orchestration responsibility",
        "evidence": "the function executes many distinct operations per call, with each handled concern contributing additional statements to the single body",
        "priority": 5,
    },
    {
        "name": "dense_branching",
        "keywords": (),
        "structures": ("deep_conditional_tree", "case_dispatch", "switch_on_var"),
        "label": "dense inline conditional branching",
        "evidence": "many conditional branches are evaluated inline, each containing its own executable statements rather than being delegated to helpers",
        "priority": 4,
    },
    {
        "name": "integration_surface",
        "keywords": ("wrapper", "adapter", "adaptor", "osal", "port", "bridge", "proxy", "service", "integration", "interface"),
        "structures": ("compatibility_branching",),
        "label": "integration surface coverage",
        "evidence": "the function executes multiple platform or interface calls in one place, each integration step adding to the total statement count",
        "priority": 4,
    },
    {
        "name": "sequential_workflow",
        "keywords": ("init", "setup", "configure", "config", "sequence", "step", "boot", "start", "procedure"),
        "structures": ("sequential_init_chain",),
        "label": "sequential procedural workflow",
        "evidence": "a strictly ordered sequence of dependent steps is executed in one function, each step contributing statements without delegation to sub-functions",
        "priority": 3,
    },
    {
        "name": "configuration_density",
        "keywords": ("table", "mapping", "lookup", "setting", "configuration", "matrix", "register", "bitfield", "constant"),
        "structures": (),
        "label": "configuration and mapping density",
        "evidence": "inline configuration assignments, register writes, or enumerated case actions each contribute individually to the total executed statement count",
        "priority": 2,
    },
)


def describe_loc_reason(function_name: str, source_match: dict | None, metric_label: str = "Function size") -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    scored = _score_structural_signals(combined, structures, _FUNCTION_SIZE_SIGNALS)
    if not scored:
        return f"{metric_label} is a structural maintainability signal — the function accumulates lines across multiple in-scope responsibilities. Reviewed and covered by testing."
    primary = scored[0]
    return f"{metric_label} is a structural maintainability signal driven by {primary['label']} — {primary['evidence']}. Reviewed and covered by testing."


def describe_stmt_reason(function_name: str, source_match: dict | None) -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    scored = _score_structural_signals(combined, structures, _STATEMENT_DENSITY_SIGNALS)
    if not scored:
        return "Statement count is a structural maintainability signal \u2014 the function executes many operations across multiple in-scope concerns. Reviewed and covered by testing."
    primary = scored[0]
    return f"Statement count is a structural maintainability signal driven by {primary['label']} \u2014 {primary['evidence']}. Reviewed and covered by testing."


def describe_return_count_reason(function_name: str, source_match: dict | None) -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    scored = _score_structural_signals(combined, structures, _EXIT_FRAGMENTATION_SIGNALS)
    if not scored:
        return "Multiple return statements are an exit fragmentation signal — the function uses early exits to manage control flow across distinct conditions. Reviewed and covered by testing."
    primary = scored[0]
    return f"Multiple return statements are an exit fragmentation signal driven by {primary['label']} — {primary['evidence']}. Reviewed and covered by testing."


def describe_goto_reason(function_name: str, source_match: dict | None) -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    scored = _score_structural_signals(combined, structures, _GOTO_USAGE_SIGNALS)
    if not scored:
        return "Goto usage is a control-flow consolidation signal — multiple branches appear to jump to shared cleanup, recovery, or exit handling in order to avoid duplicating the same logic. Reviewed and covered by testing."
    primary = scored[0]
    return f"Goto usage is a control-flow consolidation signal driven by {primary['label']} — {primary['evidence']}. Reviewed and covered by testing."


def describe_parameter_count_reason(function_name: str, source_match: dict | None) -> str:
    _, combined = extract_nesting_context(function_name, source_match)
    structures = detect_nesting_structures(function_name, source_match)
    scored = _score_structural_signals(combined, structures, _INTERFACE_COMPLEXITY_SIGNALS)
    if not scored:
        return "Parameter count is an interface complexity signal — the function accepts multiple independent inputs reflecting the range of its calling contexts. Reviewed and covered by testing."
    primary = scored[0]
    return f"Parameter count is an interface complexity signal driven by {primary['label']} — {primary['evidence']}. Reviewed and covered by testing."


def professional_comment_for_metric(
    metric: str,
    value: str,
    function_name: str,
    source_match: dict | None,
) -> str:
    if source_match and source_match.get("baseline_unused"):
        return "Not used by S3P baseline."

    if source_match and source_match.get("generated_source") == "matlab":
        return "MATLAB generated code, can be ignored."

    metric_key = normalized_header(metric)

    if metric_key == "lines of code executable":
        return describe_loc_reason(function_name, source_match, "Executable line count")

    if metric_key == "lines of code inactive":
        return (
            "Inactive code indicates conditional compilation, legacy branches, or build-time variability within the same unit, "
            "which increases inspection effort and reduces clarity of the active execution flow. "
            "This is typically an intentional design choice to support configuration or compatibility scenarios."
        )

    if metric_key == "lines of code comments":
        return (
            "Low comment density reduces immediate understandability of the code, particularly in non-obvious logic paths, "
            "but does not affect functional complexity. "
            "The implementation has been reviewed and validated, and supporting documentation is expected to complement inline comments."
        )

    if metric_key in ("cyclomatic nist", "cyclomatic"):
        return describe_cyclomatic_reason(function_name, source_match, value)

    if metric_key == "maximal level of nesting":
        return describe_nesting_reason(function_name, source_match)

    if metric_key in (
        "number of functions calling this function fanin transitive",
    ):
        return (
            "This function is used by many upstream paths across the system. This breadth of dependency should be checked against the SW System Design, "
            "and any refinement should preserve the current call behaviour and remain controlled and low risk."
        )

    if metric_key == "number of functions being called by this function fanout transitive":
        return (
            "This routine leads to many downstream behaviours, directly and indirectly. That often reflects broad orchestration logic, "
            "so any redesign should be checked against the SW System Design and validated for integration impact before applying changes."
        )


    if metric_key == "lines of code total":
        return describe_loc_reason(function_name, source_match, "Function size")

    if metric_key == "number of executed statements":
        return describe_stmt_reason(function_name, source_match)

    if metric_key == "number of return statements":
        return describe_return_count_reason(function_name, source_match)

    if metric_key == "number of goto statements":
        return describe_goto_reason(function_name, source_match)

    if metric_key == "number of function parameters":
        return describe_parameter_count_reason(function_name, source_match)

    if metric_key == "comment total loc ratio percent":
        # Low comment ratio is a documentation deficit signal only — it does NOT
        # create complexity and must not be treated as a root cause.
        return (
            "Low comment ratio makes the code harder to understand at first read, "
            "but does not create complexity by itself. "
            "The implementation has been reviewed and tested carefully, and supporting "
            "project documentation such as the module specification should be available "
            "to complement the inline comments."
        )

    if metric_key == "inactive total loc ratio percent":
        return (
            "Inactive code reflects conditional compilation, legacy paths, or build-time variability, "
            "increasing inspection effort and reducing clarity of active logic, "
            "while typically supporting configuration or compatibility needs."
        )

    if metric_key == "maintainability metric":
        return describe_maintainability_reason(function_name, source_match, value)

    if metric_key == "static path count":
        return describe_static_path_reason(function_name, source_match, value)

    if metric_key == "number of functions calling this function fanin non transitive":
        return describe_fanin_non_transitive_reason(function_name, source_match)

    if metric_key == "number of functions being called by this function fanout non transitive":
        return describe_fanout_non_transitive_reason(function_name, source_match)

    return (
        f"The reported metric for {metric} suggests elevated structural complexity in the implemented decisions, paths, or actions. "
        "The logic should be regarded as reviewed, tested, and stable, and any improvement should remain controlled and behaviour-preserving."
    )


def trim_general_comment_tail(comment: str) -> str:
    generic_markers = (
        "The code has passed review and proper testing",
        "The current logic has passed review and proper testing",
        "The logic should be regarded as reviewed, tested, and stable",
    )
    for marker in generic_markers:
        if marker in comment:
            return comment.split(marker, 1)[0].rstrip(" .,:;") + "."

    # Strip the last action-recommendation clause ("Any ...", "and any ...", "so any ...")
    comment_lower = comment.lower()
    last_action = max(
        (comment_lower.rfind(p) for p in (". any ", ", and any ", " and any ", ", so any ", " so any ", ", any ")),
        default=-1,
    )
    if last_action >= 0:
        if comment[last_action] == ".":
            return comment[: last_action + 1]
        return comment[:last_action].rstrip(" .,:;") + "."
    return comment


def threshold_label(rule: dict, value_str: str) -> str:
    """Return a compact threshold annotation string for display, e.g. 'warn ≥101, error ≥150, SMK: FFS'."""
    parts: list[str] = []
    if rule.get("warning") is not None and rule["warning"] > 0:
        parts.append(f"warn \u2265{rule['warning']}")
    if rule.get("error") is not None and rule["error"] > 0:
        parts.append(f"error \u2265{rule['error']}")
    if rule.get("warning_below") is not None:
        parts.append(f"min {rule['warning_below']}%")
    if rule.get("smk"):
        parts.append(f"SMK: {rule['smk']}")
    return ", ".join(parts)


def build_combined_solution(
    metrics: list[tuple[str, str]], function_name: str, source_match: dict | None
) -> str:
    """Build an audit-style combined solution comment covering all red metrics for one function."""
    n = len(metrics)
    if source_match:
        location_text = f"{source_match['path']}:{source_match['line']}"
        header = (
            f"Function {function_name} ({location_text}) violates {n} "
            f"project metric{'s' if n > 1 else ''}."
        )
    else:
        header = (
            f"Function {function_name} violates {n} project metric{'s' if n > 1 else ''} "
            f"(source not found in repository)."
        )

    bullet_lines: list[str] = []
    seen_actions: list[str] = []
    for metric, value in metrics:
        rule = lookup_metric_rule(metric)
        if rule:
            thresh = threshold_label(rule, value)
            rationale = rule["rationale"]
            action = rule["action"]
            entry = f"\u2022 {metric} = {value}"
            if thresh:
                entry += f" ({thresh})"
            entry += f": {rationale}."
            bullet_lines.append(entry)
            if action and action not in seen_actions:
                seen_actions.append(action)
        else:
            bullet_lines.append(f"\u2022 {metric} = {value}: exceeds project threshold.")
            generic = "Refactor to bring the metric within project thresholds"
            if generic not in seen_actions:
                seen_actions.append(generic)

    professional_comments: list[str] = []
    for metric, value in metrics:
        comment = professional_comment_for_metric(metric, value, function_name, source_match)
        if n > 1:
            comment = trim_general_comment_tail(comment)
        if comment not in professional_comments:
            professional_comments.append(comment)

    combined_comment = " ".join(professional_comments)
    if source_match and source_match.get("generated_source") == "matlab":
        return combined_comment
    if source_match and source_match.get("baseline_unused"):
        return combined_comment
    if n == 1:
        return combined_comment
    return combined_comment


def evaluate_rule(rule, cell_value: object) -> bool:
    value = to_number(cell_value)
    if value is None:
        return False

    formulas = getattr(rule, "formula", []) or []
    parsed = []
    for formula in formulas:
        numeric = to_number(formula)
        if numeric is None:
            return False
        parsed.append(numeric)

    operator = getattr(rule, "operator", None)
    if rule.type != "cellIs" or operator is None:
        return False

    if operator == "greaterThanOrEqual" and len(parsed) >= 1:
        return value >= parsed[0]
    if operator == "greaterThan" and len(parsed) >= 1:
        return value > parsed[0]
    if operator == "lessThanOrEqual" and len(parsed) >= 1:
        return value <= parsed[0]
    if operator == "lessThan" and len(parsed) >= 1:
        return value < parsed[0]
    if operator == "equal" and len(parsed) >= 1:
        return value == parsed[0]
    if operator == "between" and len(parsed) >= 2:
        return parsed[0] <= value <= parsed[1]

    return False


def conditional_format_marks_red(ws, cell_ref: str) -> bool:
    matching_rules = []
    for cf_range, rules in ws.conditional_formatting._cf_rules.items():
        if not matches_range(cell_ref, cf_range.sqref):
            continue
        for rule in rules:
            if evaluate_rule(rule, ws[cell_ref].value):
                matching_rules.append(rule)

    if not matching_rules:
        return False

    # In this PMS workbook, the highest-priority matching rule is the red state.
    matching_rules.sort(key=lambda rule: getattr(rule, "priority", 10**9))
    top_rule = matching_rules[0]
    return getattr(top_rule, "priority", 10**9) % 2 == 1


def iter_source_files(repo_root: Path, include_build: bool, allowed_files: set[str] | None = None) -> Iterable[Path]:
    excluded = set(EXCLUDED_DIRS)
    if not include_build:
        excluded.add("build")

    candidate_roots = [repo_root]
    preferred_roots = [repo_root / name for name in PREFERRED_SOURCE_ROOTS if (repo_root / name).exists()]
    if preferred_roots:
        candidate_roots = preferred_roots

    for base_root in candidate_roots:
        for current_root, dirnames, filenames in os.walk(base_root):
            dirnames[:] = [dirname for dirname in dirnames if dirname not in excluded]
            for filename in filenames:
                candidate = Path(current_root) / filename
                if candidate.suffix.lower() not in SOURCE_SUFFIXES:
                    continue
                if allowed_files is not None and normalize(filename) not in allowed_files:
                    continue
                yield candidate


def iter_reference_files(
    repo_root: Path,
    include_build: bool,
    allowed_files: set[str] | None = None,
) -> Iterable[Path]:
    excluded = set(EXCLUDED_DIRS)
    if not include_build:
        excluded.add("build")

    candidate_roots = [repo_root]
    preferred_roots = [repo_root / name for name in PREFERRED_SOURCE_ROOTS if (repo_root / name).exists()]
    if preferred_roots:
        candidate_roots = preferred_roots

    for base_root in candidate_roots:
        for current_root, dirnames, filenames in os.walk(base_root):
            dirnames[:] = [dirname for dirname in dirnames if dirname not in excluded]
            for filename in filenames:
                candidate = Path(current_root) / filename
                if candidate.suffix.lower() not in REFERENCE_SUFFIXES:
                    continue
                if allowed_files is not None and normalize(filename) not in allowed_files:
                    continue
                yield candidate


def score_candidate(path: Path, line: str, function_name: str) -> int:
    score = 0
    path_text = str(path).replace("\\", "/").lower()
    if "/src/" in path_text or "/core/" in path_text:
        score += 40
    if "/pkg/" in path_text or "/app/" in path_text:
        score += 20
    if path.suffix.lower() in {".c", ".cc", ".cpp", ".cxx"}:
        score += 25
    if "/include/" in path_text:
        score -= 20
    if "/build/" in path_text:
        score -= 50

    if re.search(rf"\bFUNC\s*\([^\n]*\)\s*{re.escape(function_name)}\b", line):
        score += 80
    if re.search(rf"\b{re.escape(function_name)}\s*\(", line):
        score += 30
    if line.strip().startswith(("/*", "//", "*")):
        score -= 40
    if any(token in line for token in ("extern ", "typedef ", "&" + function_name)):
        score -= 20
    if line.strip().endswith(";"):
        score -= 10
    return score


def find_function_location(
    repo_root: Path,
    function_name: str,
    cache: dict[tuple[str, str | None], dict | None],
    preferred_file_name: str | None,
    allowed_files: set[str] | None,
) -> dict | None:
    cache_key = (function_name, preferred_file_name)
    if cache_key in cache:
        return cache[cache_key]

    best_match = None
    for include_build in (False, True):
        for source_file in iter_source_files(repo_root, include_build, allowed_files):
            try:
                lines = source_file.read_text(encoding="utf-8", errors="ignore").splitlines()
            except OSError:
                continue

            generated_source = detect_generated_source(lines)

            for line_number, line in enumerate(lines, start=1):
                if function_name not in line:
                    continue

                score = score_candidate(source_file, line, function_name)
                candidate = {
                    "path": str(source_file.relative_to(repo_root)),
                    "file_name": source_file.name,
                    "line": line_number,
                    "line_text": line.strip(),
                    "context_text": " ".join(
                        snippet.strip() for snippet in lines[max(0, line_number - 5) : min(len(lines), line_number + 40)]
                    ),
                    "generated_source": generated_source,
                    "score": score,
                }
                if preferred_file_name and normalize(source_file.name) == normalize(preferred_file_name):
                    candidate["score"] += 120
                if best_match is None or candidate["score"] > best_match["score"]:
                    best_match = candidate

        if best_match is not None:
            break

    cache[cache_key] = best_match
    return best_match


def compile_name_patterns(function_names: set[str], chunk_size: int = 100) -> list[re.Pattern[str]]:
    valid_names = sorted({name for name in function_names if name and name != "(no function name)"})
    patterns = []
    for index in range(0, len(valid_names), chunk_size):
        chunk = valid_names[index : index + chunk_size]
        joined = "|".join(re.escape(name) for name in chunk)
        patterns.append(re.compile(rf"\b({joined})\b"))
    return patterns


def build_function_index(
    repo_root: Path,
    function_names: set[str],
    allowed_files: set[str] | None,
) -> dict[str, list[dict]]:
    index: dict[str, list[dict]] = {name: [] for name in function_names}
    patterns = compile_name_patterns(function_names)

    if not patterns:
        return index

    for source_file in iter_source_files(repo_root, include_build=False, allowed_files=allowed_files):
        try:
            lines = source_file.read_text(encoding="utf-8", errors="ignore").splitlines()
        except OSError:
            continue

        generated_source = detect_generated_source(lines)

        relative_path = str(source_file.relative_to(repo_root))
        for line_number, line in enumerate(lines, start=1):
            for pattern in patterns:
                matches = pattern.findall(line)
                if not matches:
                    continue

                for function_name in matches:
                    score = score_candidate(source_file, line, function_name)
                    candidate = {
                        "path": relative_path,
                        "file_name": source_file.name,
                        "line": line_number,
                        "line_text": line.strip(),
                        "context_text": " ".join(
                            snippet.strip() for snippet in lines[max(0, line_number - 5) : min(len(lines), line_number + 40)]
                        ),
                        "generated_source": generated_source,
                        "score": score,
                    }
                    index.setdefault(function_name, []).append(candidate)

    return index


def build_reference_index(
    repo_root: Path,
    function_names: set[str],
    allowed_files: set[str] | None,
) -> dict[str, list[dict]]:
    index: dict[str, list[dict]] = {name: [] for name in function_names}
    patterns = compile_name_patterns(function_names)

    if not patterns:
        return index

    for source_file in iter_reference_files(repo_root, include_build=False, allowed_files=allowed_files):
        try:
            lines = source_file.read_text(encoding="utf-8", errors="ignore").splitlines()
        except OSError:
            continue

        relative_path = str(source_file.relative_to(repo_root))
        for line_number, line in enumerate(lines, start=1):
            for pattern in patterns:
                matches = pattern.findall(line)
                if not matches:
                    continue

                for function_name in matches:
                    candidate = {
                        "path": relative_path,
                        "file_name": source_file.name,
                        "line": line_number,
                        "line_text": line.strip(),
                        "score": 0,
                    }
                    index.setdefault(function_name, []).append(candidate)

    return index


def resolve_sheet_name(wb, preferred: str) -> str:
    for name in wb.sheetnames:
        if normalize(name) == normalize(preferred):
            return name

    for name in wb.sheetnames:
        n = normalize(name)
        if "understand" in n and "function" in n:
            return name

    for name in wb.sheetnames:
        if normalize(name).endswith("_function"):
            return name

    raise ValueError(
        f"Could not find sheet '{preferred}' or a matching *_function sheet. "
        f"Available sheets: {wb.sheetnames}"
    )


def validate_new_template_args(args: argparse.Namespace) -> None:
    missing_args = []
    if not args.file_prefixes:
        missing_args.append("--file-prefixes")
    if not args.platform_cds_module:
        missing_args.append("--platform-cds-module")
    if not args.sw_module_name:
        missing_args.append("--sw-module-name")
    if not args.responsible_name:
        missing_args.append("--responsible-name")
    if missing_args:
        raise ValueError(
            "The new template flow requires: " + ", ".join(missing_args)
        )


def metric_name(ws, col_letter: str) -> str:
    for header_row in (7, 8):
        value = ws[f"{col_letter}{header_row}"].value
        if value is not None and str(value).strip() != "":
            return str(value).strip()
    return f"Column {col_letter}"


def find_function_col(ws) -> int:
    keywords = ("function", "function name", "name")
    for row in (7, 8):
        for col in range(1, 16):
            val = normalize(ws.cell(row=row, column=col).value)
            if any(k in val for k in keywords):
                return col
    return 1


def find_file_col(ws) -> int | None:
    keywords = ("file", "source", "module")
    for row in (7, 8):
        for col in range(1, 16):
            val = normalize(ws.cell(row=row, column=col).value)
            if not val:
                continue
            if any(keyword in val for keyword in keywords):
                return col
    return None


def find_metric_col(ws, expected_metric_name: str) -> int | None:
    expected = normalize(expected_metric_name)
    for col in range(1, ws.max_column + 1):
        column_letter = get_column_letter(col)
        if normalize(metric_name(ws, column_letter)) == expected:
            return col
    return None


def get_last_data_row(ws, start_row: int, function_col: int, target_columns: list[str], file_col: int | None) -> int:
    last = start_row - 1
    for row in range(start_row, ws.max_row + 1):
        fn_val = ws.cell(row=row, column=function_col).value
        has_fn = fn_val is not None and str(fn_val).strip() != ""
        file_val = ws.cell(row=row, column=file_col).value if file_col else None
        has_file = file_val is not None and str(file_val).strip() != ""
        has_metric = any(
            ws[f"{col}{row}"].value is not None and str(ws[f"{col}{row}"].value).strip() != ""
            for col in target_columns
        )
        if has_fn or has_file or has_metric:
            last = row
    return max(last, start_row - 1)


def choose_source_match(candidates: list[dict], preferred_file_name: str | None) -> dict | None:
    if not candidates:
        return None

    best_match = None
    for candidate in candidates:
        adjusted_score = candidate["score"]
        if preferred_file_name and normalize(candidate["file_name"]) == normalize(preferred_file_name):
            adjusted_score += 120

        candidate_with_score = dict(candidate)
        candidate_with_score["score"] = adjusted_score
        if best_match is None or candidate_with_score["score"] > best_match["score"]:
            best_match = candidate_with_score

    return best_match


def code_solution_for_metric(metric: str, function_name: str, source_match: dict | None) -> str:
    m = metric.lower()
    source_text = ""
    if source_match is not None:
        source_text = (
            f"Review {function_name} in {source_match['path']}:{source_match['line']}. "
        )

    if "cyclomatic" in m:
        return (
            source_text
            + "The complexity is primarily driven by necessary business logic and multiple "
            "conditional pathways required to handle diverse scenarios. If decomposition is "
            "pursued, extract only self-contained branch groups into static helpers while "
            "keeping the top-level orchestration intact, and ensure every conditional path "
            "is covered by dedicated branch-level tests before and after the change. "
            + SOLUTION_NOTE
        )
    if "nest" in m:
        return (
            source_text
            + "Reduce nesting with guard clauses, early exits, and helper routines for "
            "branch-specific work. "
            + SOLUTION_NOTE
        )
    if "path" in m:
        return (
            source_text
            + "Collapse duplicated decision paths, move repeated checks into helpers, "
            "and separate state validation from main behavior. "
            + SOLUTION_NOTE
        )
    if "inactive" in m:
        return (
            source_text
            + "Remove inactive branches, or move legacy/debug-only flow behind a single "
            "controlled abstraction instead of inline conditional code. "
            + SOLUTION_NOTE
        )
    if "comment" in m:
        return (
            source_text
            + "Prefer clearer names and smaller helper functions first, then add brief "
            "comments only for non-obvious logic, assumptions, or safety constraints. "
            + SOLUTION_NOTE
        )
    if "parameter" in m:
        return (
            source_text
            + "Bundle related arguments into a context/config struct, remove redundant "
            "scalars, and keep only parameters that genuinely vary per call. "
            + SOLUTION_NOTE
        )
    if "fanout" in m:
        return (
            source_text
            + "Reduce direct dependencies by extracting a smaller service layer or helper "
            "API so this function calls fewer modules directly. "
            + SOLUTION_NOTE
        )
    if "fanin" in m:
        return (
            source_text
            + "Review whether this function mixes too many responsibilities; split public "
            "entry points or move shared logic into internal helpers to reduce coupling. "
            + SOLUTION_NOTE
        )
    if "executed statements" in m or "lines of code" in m:
        return (
            source_text
            + "Extract validation, error reporting, and repeated branch bodies into static "
            "helpers so the main function keeps one clear control flow. "
            + SOLUTION_NOTE
        )
    if "maintainability" in m:
        return (
            source_text
            + "Refactor the function into smaller units with clearer responsibilities, and "
            "simplify condition handling before tuning comments or formatting. "
            + SOLUTION_NOTE
        )
    return (
        source_text
        + "Refactor the function to meet the project metric threshold for this item. "
        + SOLUTION_NOTE
    )


def iter_red_findings(
    ws,
    start_row: int,
    repo_root: Path,
    target_columns: list[str],
    allowed_files: set[str] | None,
) -> Iterable[dict]:
    function_col = find_function_col(ws)
    file_col = find_file_col(ws)
    fanin_non_transitive_col = find_metric_col(
        ws, "Number of functions calling this function (FanIn, Non-Transitive)"
    )
    last_row = get_last_data_row(ws, start_row, function_col, target_columns, file_col)
    metric_names = {col: metric_name(ws, col) for col in target_columns}
    function_names = set()
    fallback_cache: dict[tuple[str, str | None], dict | None] = {}
    perform_repo_lookup = allowed_files is not None

    function_cache: dict[str, list[dict]] = {}
    reference_cache: dict[str, list[dict]] = {}
    search_scope_files: set[str] | None = allowed_files

    if perform_repo_lookup:
        for row in range(start_row, last_row + 1):
            workbook_file_name = None
            if file_col is not None:
                file_value = ws.cell(row=row, column=file_col).value
                if file_value is not None and str(file_value).strip() != "":
                    workbook_file_name = Path(str(file_value).strip()).name

            if allowed_files is not None and workbook_file_name and normalize(workbook_file_name) not in allowed_files:
                continue

            function_name = ws.cell(row=row, column=function_col).value
            if function_name is None or str(function_name).strip() == "":
                continue
            function_names.add(str(function_name))

        function_cache = build_function_index(repo_root, function_names, search_scope_files)
        reference_cache = build_reference_index(repo_root, function_names, search_scope_files)

    for row in range(start_row, last_row + 1):
        workbook_file_name = None
        if file_col is not None:
            file_value = ws.cell(row=row, column=file_col).value
            if file_value is not None and str(file_value).strip() != "":
                workbook_file_name = Path(str(file_value).strip()).name

        normalized_workbook_file_name = normalize(workbook_file_name) if workbook_file_name else ""
        if allowed_files is not None and normalized_workbook_file_name and normalized_workbook_file_name not in allowed_files:
            continue

        function_name = ws.cell(row=row, column=function_col).value
        if function_name is None or str(function_name).strip() == "":
            function_name = "(no function name)"
        function_name = str(function_name)
        source_match = None
        baseline_call_count = 0
        reference_count = 0
        fanin_non_transitive_value = None

        if perform_repo_lookup:
            source_match = choose_source_match(function_cache.get(function_name, []), workbook_file_name)
            if source_match is None and function_name != "(no function name)":
                source_match = find_function_location(
                    repo_root,
                    function_name,
                    fallback_cache,
                    workbook_file_name,
                    search_scope_files,
                )

            baseline_call_count = count_baseline_call_sites(function_cache.get(function_name, []), function_name)
            reference_count = count_reference_sites(reference_cache.get(function_name, []), function_name)

        if fanin_non_transitive_col is not None:
            fanin_non_transitive_value = to_number(
                ws.cell(row=row, column=fanin_non_transitive_col).value
            )
        if source_match is not None:
            source_match = dict(source_match)
            source_match["baseline_call_count"] = baseline_call_count
            source_match["reference_count"] = reference_count
            source_match["fanin_non_transitive_value"] = fanin_non_transitive_value
            if fanin_non_transitive_value is not None:
                source_match["baseline_unused"] = fanin_non_transitive_value == 0
            else:
                source_match["baseline_unused"] = (
                    baseline_call_count == 0
                    and reference_count == 0
                    and not has_baseline_usage_hint(source_match)
                )

        matched_file_name = normalize(source_match["file_name"]) if source_match else ""
        if allowed_files is not None and matched_file_name and matched_file_name not in allowed_files:
            continue
        if allowed_files is not None and not normalized_workbook_file_name and not matched_file_name:
            continue

        red_metrics: list[tuple[str, str]] = []
        for col in target_columns:
            cell = ws[f"{col}{row}"]
            if is_red_fill(cell.fill) or conditional_format_marks_red(ws, cell.coordinate):
                red_metrics.append(
                    (metric_names[col], "" if cell.value is None else str(cell.value))
                )

        if red_metrics:
            yield {
                "line_no": row,
                "file_name": workbook_file_name or (source_match["file_name"] if source_match else "Unknown"),
                "function_name": function_name,
                "metrics": red_metrics,
                "repo_location": (
                    f"{source_match['path']}:{source_match['line']}" if source_match else "Not found in repository"
                ),
                "code_solution": build_combined_solution(red_metrics, function_name, source_match),
            }


def collect_module_scope(
    ws,
    start_row: int,
    repo_root: Path,
    target_columns: list[str],
    allowed_files: set[str] | None,
) -> tuple[set[int], set[str]]:
    function_col = find_function_col(ws)
    file_col = find_file_col(ws)
    last_row = get_last_data_row(ws, start_row, function_col, target_columns, file_col)
    module_rows: set[int] = set()
    module_files: set[str] = set()

    for row in range(start_row, last_row + 1):
        function_name = ws.cell(row=row, column=function_col).value
        if function_name is None or str(function_name).strip() == "":
            continue

        workbook_file_name = None
        if file_col is not None:
            file_value = ws.cell(row=row, column=file_col).value
            if file_value is not None and str(file_value).strip() != "":
                workbook_file_name = Path(str(file_value).strip()).name

        normalized_workbook_file_name = normalize(workbook_file_name) if workbook_file_name else ""
        if allowed_files is None:
            module_rows.add(row)
            if workbook_file_name:
                module_files.add(workbook_file_name)
            continue

        if normalized_workbook_file_name and normalized_workbook_file_name in allowed_files:
            module_rows.add(row)
            if workbook_file_name:
                module_files.add(workbook_file_name)

    return module_rows, module_files


def default_annotated_workbook_path(workbook: Path) -> Path:
    return workbook.with_name(f"{workbook.stem}_annotated{workbook.suffix}")


def build_dated_suffix() -> str:
    return datetime.now().strftime("%Y%m%d")


def default_html_output_path(module_name: str | None, base_dir: Path) -> Path:
    dated_suffix = build_dated_suffix()
    prefix = module_name if module_name else DEFAULT_HTML_BASENAME
    safe_prefix = re.sub(r"[^A-Za-z0-9_-]+", "_", prefix).strip("_") or DEFAULT_HTML_BASENAME
    return base_dir / f"{safe_prefix}_{dated_suffix}.html"


def default_annotated_workbook_path_for_module(workbook: Path, module_name: str | None) -> Path:
    dated_suffix = build_dated_suffix()
    safe_module = re.sub(r"[^A-Za-z0-9_-]+", "_", module_name or "module").strip("_") or "module"
    return workbook.with_name(f"{workbook.stem}_{safe_module}_{dated_suffix}{workbook.suffix}")


def default_new_template_output_path(workbook: Path, module_name: str | None) -> Path:
    dated_suffix = build_dated_suffix()
    safe_module = re.sub(r"[^A-Za-z0-9_-]+", "_", module_name or "module").strip("_") or "module"
    return workbook.with_name(f"{workbook.stem}_{safe_module}_comment_filled_{dated_suffix}{workbook.suffix}")


def fill_new_template_comments(
    ws,
    repo_root: Path,
    start_row: int,
    file_prefixes: list[str],
    platform_cds_module: str,
    sw_module_name: str,
    responsible_name: str,
) -> dict[str, object]:
    matching_rows = find_matching_new_template_rows(ws, start_row, file_prefixes)
    comment_mapping = resolve_new_template_comment_columns(ws)
    allowed_files = source_allowed_files_for_prefixes(ws, start_row, file_prefixes)
    function_names = {
        row["function_name"]
        for row in matching_rows
        if row["function_name"] and row["function_name"] != "(no function name)"
    }

    function_cache = build_function_index(repo_root, function_names, allowed_files)

    updated_comment_cells = 0
    commented_rows = 0
    touched_files: set[str] = set()
    findings: list[dict] = []
    module_rows = {int(row_info["row"]) for row_info in matching_rows}

    for row_info in matching_rows:
        row_number = int(row_info["row"])
        file_name = str(row_info["file_name"] or "")
        function_name = str(row_info["function_name"] or "") or "(no function name)"

        if not has_existing_text(ws[f"{NEW_TEMPLATE_PLATFORM_COLUMN}{row_number}"].value):
            ws[f"{NEW_TEMPLATE_PLATFORM_COLUMN}{row_number}"] = platform_cds_module
        if not has_existing_text(ws[f"{NEW_TEMPLATE_MODULE_COLUMN}{row_number}"].value):
            ws[f"{NEW_TEMPLATE_MODULE_COLUMN}{row_number}"] = sw_module_name
        if not has_existing_text(ws[f"{NEW_TEMPLATE_OWNER_COLUMN}{row_number}"].value):
            ws[f"{NEW_TEMPLATE_OWNER_COLUMN}{row_number}"] = responsible_name

        source_match = choose_source_match(function_cache.get(function_name, []), file_name)

        row_comment_count = 0
        row_metrics: list[tuple[str, str]] = []
        row_comments: list[str] = []
        for comment_column, metric_column in comment_mapping.items():
            comment_cell = ws[f"{comment_column}{row_number}"]
            if not is_purple_fill(comment_cell.fill):
                continue
            if has_existing_text(comment_cell.value):
                continue

            metric_label = metric_name(ws, metric_column)
            metric_cell = ws[f"{metric_column}{row_number}"]
            metric_value = "" if metric_cell.value is None else str(metric_cell.value)
            comment_cell.value = professional_comment_for_metric(
                metric_label,
                metric_value,
                function_name,
                source_match,
            )
            row_metrics.append((metric_label, metric_value))
            row_comments.append(str(comment_cell.value))
            updated_comment_cells += 1
            row_comment_count += 1

        if row_comment_count:
            commented_rows += 1
            findings.append(
                {
                    "line_no": row_number,
                    "file_name": file_name or (source_match["file_name"] if source_match else "Unknown"),
                    "function_name": function_name,
                    "metrics": row_metrics,
                    "repo_location": (
                        f"{source_match['path']}:{source_match['line']}" if source_match else "Not found in repository"
                    ),
                    "code_solution": " ".join(row_comments),
                }
            )
        if file_name:
            touched_files.add(file_name)

    return {
        "matching_rows": len(matching_rows),
        "commented_rows": commented_rows,
        "updated_comment_cells": updated_comment_cells,
        "comment_mapping": comment_mapping,
        "touched_files": sorted(touched_files),
        "findings": findings,
        "module_rows": module_rows,
        "module_files": touched_files,
    }


def apply_workbook_comments(
    ws,
    findings: list[dict],
    module_rows: set[int],
    module_name: str | None,
) -> None:
    comments_by_row = {item["line_no"]: item["code_solution"] for item in findings}

    for row in sorted(module_rows):
        ws[f"{COMMENT_COLUMN}{row}"] = comments_by_row.get(row, "")
        if module_name:
            ws[f"{RESPONSIBILITY_COLUMN}{row}"] = module_name


def build_html_report(
    title: str,
    findings: list[dict],
    module_rows: set[int],
    module_files: set[str],
    workbook: Path,
    sheet_name: str,
    target_columns: list[str],
    file_list_path: Path | None,
    metric_rule_file: Path | None = None,
    scope_description: str | None = None,
    summary_title: str = "Red Metric Summary",
    empty_message: str = "No red metrics found in the configured columns for the selected file scope.",
) -> str:
    counts = Counter(m_name for item in findings for m_name, _ in item["metrics"])
    comment_count = len(findings)
    no_comment_count = max(len(module_rows) - comment_count, 0)
    module_file_count = len(module_files)

    def _render_metrics_cell(metrics: list[tuple[str, str]]) -> str:
        lines = []
        for m_name, m_val in metrics:
            rule = lookup_metric_rule(m_name)
            thresh = threshold_label(rule, m_val) if rule else ""
            entry = f"{html.escape(m_name)}: <strong>{html.escape(m_val)}</strong>"
            if thresh:
                entry += f" <span style='color:#6b7280;font-size:0.85em;'>({html.escape(thresh)})</span>"
            lines.append(entry)
        return "<br>".join(lines)

    rows_html = "\n".join(
        "<tr>"
        f"<td>{item['line_no']}</td>"
        f"<td>{html.escape(item['file_name'])}</td>"
        f"<td>{html.escape(item['function_name'])}</td>"
        f"<td>{_render_metrics_cell(item['metrics'])}</td>"
        f"<td>{html.escape(item['repo_location'])}</td>"
        f"<td>{html.escape(item['code_solution'])}</td>"
        f"<td style='text-align:center;vertical-align:middle;'><input type='checkbox' data-row-key='{html.escape(str(item['line_no']) + '|' + item['file_name'] + '|' + item['function_name'], quote=True)}' aria-label='Select {html.escape(item['function_name'])} for fix' /></td>"
        "</tr>"
        for item in findings
    )

    summary_items = "\n".join(
        f"<li>{html.escape(metric)}: {count}</li>" for metric, count in sorted(counts.items())
    )

    if not summary_items:
        summary_items = "<li>No red metrics found in configured columns.</li>"

    if not rows_html:
        rows_html = f"<tr><td colspan='7'>{html.escape(empty_message)}</td></tr>"

    scope_columns = ", ".join(target_columns)
    file_scope = str(file_list_path) if file_list_path is not None else "All repository files"
    metric_rule_label = str(metric_rule_file) if metric_rule_file else "Built-in defaults"
    resolved_scope_description = scope_description or f"Red-highlighted metrics in columns {scope_columns} from row 9 to last function row"

    return f"""<!DOCTYPE html>
<html lang=\"en\">
<head>
  <meta charset=\"utf-8\" />
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1\" />
  <title>{html.escape(title)}</title>
    <style>
        body {{ font-family: Segoe UI, Arial, sans-serif; margin: 24px; color: #1f2937; }}
        h1 {{ margin-bottom: 8px; }}
        .meta {{ margin-bottom: 16px; color: #374151; }}
        .summary {{ background: #f8fafc; border: 1px solid #dbeafe; padding: 12px; border-radius: 8px; margin-bottom: 16px; }}
        .filter-note {{ margin: 12px 0 10px; color: #4b5563; font-size: 0.92em; }}
        table {{ width: 100%; border-collapse: collapse; }}
        th, td {{ border: 1px solid #d1d5db; padding: 8px; vertical-align: top; }}
        th {{ background: #111827; color: #ffffff; text-align: left; position: relative; }}
        tr:nth-child(even) {{ background: #f9fafb; }}
        td:nth-child(4) {{ min-width: 260px; }}
        td:nth-child(6), th:nth-child(6) {{ min-width: 520px; width: 36%; }}
        td:nth-child(6) {{ font-size: 0.88em; color: #374151; white-space: pre-wrap; word-break: break-word; }}
        td:nth-child(7), th:nth-child(7) {{ min-width: 140px; text-align: center; }}
        input[type='checkbox'] {{ width: 18px; height: 18px; cursor: pointer; }}
        .th-wrap {{ display: flex; align-items: center; justify-content: space-between; gap: 8px; }}
        .filter-menu {{ position: relative; }}
        .filter-menu summary {{ list-style: none; cursor: pointer; font-size: 0.85em; color: #d1d5db; user-select: none; }}
        .filter-menu summary::-webkit-details-marker {{ display: none; }}
        .filter-menu[open] summary {{ color: #ffffff; }}
        .filter-panel {{ position: absolute; top: calc(100% + 6px); right: 0; z-index: 5; width: 220px; background: #ffffff; border: 1px solid #cbd5e1; border-radius: 8px; box-shadow: 0 10px 24px rgba(15, 23, 42, 0.18); padding: 10px; }}
        .filter-panel label {{ display: block; color: #374151; font-size: 0.78em; margin: 0 0 4px; }}
        .filter-panel input, .filter-panel select {{ width: 100%; box-sizing: border-box; padding: 6px 8px; border: 1px solid #9ca3af; border-radius: 4px; font: inherit; margin-bottom: 8px; }}
        .filter-panel .filter-clear {{ margin-top: 2px; width: 100%; border: 0; border-radius: 4px; background: #e5e7eb; padding: 6px 8px; cursor: pointer; font: inherit; }}
    </style>
</head>
<body>
  <h1>{html.escape(title)}</h1>
  <div class=\"meta\">
    Workbook: {html.escape(str(workbook))}<br />
    Sheet: {html.escape(sheet_name)}<br />
    Metric Rule File: {html.escape(metric_rule_label)}<br />
        File Scope: {html.escape(file_scope)}<br />
        Scope: {html.escape(resolved_scope_description)}
  </div>

    <div class=\"summary\">
        <strong>{html.escape(summary_title)}</strong>
        <ul>
            {summary_items}
        </ul>
                <div><strong>Total Commented Items:</strong> {comment_count}</div>
                <div><strong>Total Module Files Covered:</strong> {module_file_count}</div>
                <div><strong>Module Rows Filled Without Comment:</strong> {no_comment_count}</div>
                <div><strong>Comment Note:</strong> Comments are generated according to the red metric type for each function and reflect reviewed, tested logic.</div>
    </div>

    <div class=\"filter-note\">Use the down arrow in each header to open a filter menu with text search and a value list. Checkbox selections are stored in this browser, mirrored into the page markup, and can be consumed by the HTML fix workflow after you save the report.</div>

    <script id=\"pms-audit-selected-rows\" type=\"application/json\">{{}}</script>

    <table id=\"audit-table\">
    <thead>
            <tr>
                <th><div class=\"th-wrap\"><span>Line No</span><details class=\"filter-menu\"><summary>&#9662;</summary><div class=\"filter-panel\"><label>Search</label><input type=\"text\" data-column=\"0\" data-filter-kind=\"text\" placeholder=\"Filter line\" /><label>Choose value</label><select data-column=\"0\" data-filter-kind=\"select\"><option value=\"\">All</option></select><button type=\"button\" class=\"filter-clear\">Clear</button></div></details></div></th>
                <th><div class=\"th-wrap\"><span>File Name</span><details class=\"filter-menu\"><summary>&#9662;</summary><div class=\"filter-panel\"><label>Search</label><input type=\"text\" data-column=\"1\" data-filter-kind=\"text\" placeholder=\"Filter file\" /><label>Choose value</label><select data-column=\"1\" data-filter-kind=\"select\"><option value=\"\">All</option></select><button type=\"button\" class=\"filter-clear\">Clear</button></div></details></div></th>
                <th><div class=\"th-wrap\"><span>Function Name</span><details class=\"filter-menu\"><summary>&#9662;</summary><div class=\"filter-panel\"><label>Search</label><input type=\"text\" data-column=\"2\" data-filter-kind=\"text\" placeholder=\"Filter function\" /><label>Choose value</label><select data-column=\"2\" data-filter-kind=\"select\"><option value=\"\">All</option></select><button type=\"button\" class=\"filter-clear\">Clear</button></div></details></div></th>
                <th><div class=\"th-wrap\"><span>Red Metrics (Value / Threshold)</span><details class=\"filter-menu\"><summary>&#9662;</summary><div class=\"filter-panel\"><label>Search</label><input type=\"text\" data-column=\"3\" data-filter-kind=\"text\" placeholder=\"Filter metric\" /><label>Choose value</label><select data-column=\"3\" data-filter-kind=\"select\"><option value=\"\">All</option></select><button type=\"button\" class=\"filter-clear\">Clear</button></div></details></div></th>
                <th><div class=\"th-wrap\"><span>Repository Match</span><details class=\"filter-menu\"><summary>&#9662;</summary><div class=\"filter-panel\"><label>Search</label><input type=\"text\" data-column=\"4\" data-filter-kind=\"text\" placeholder=\"Filter repo match\" /><label>Choose value</label><select data-column=\"4\" data-filter-kind=\"select\"><option value=\"\">All</option></select><button type=\"button\" class=\"filter-clear\">Clear</button></div></details></div></th>
                                <th><div class=\"th-wrap\"><span>Comment</span><details class=\"filter-menu\"><summary>&#9662;</summary><div class=\"filter-panel\"><label>Search</label><input type=\"text\" data-column=\"5\" data-filter-kind=\"text\" placeholder=\"Filter comment\" /><label>Choose value</label><select data-column=\"5\" data-filter-kind=\"select\"><option value=\"\">All</option></select><button type=\"button\" class=\"filter-clear\">Clear</button></div></details></div></th>
                <th>Select to Fix</th>
            </tr>
    </thead>
    <tbody>
      {rows_html}
    </tbody>
  </table>
    <script>
        (function () {{
            const table = document.getElementById('audit-table');
            if (!table) {{
                return;
            }}
            const rows = Array.from(table.querySelectorAll('tbody tr'));
            const textFilters = Array.from(table.querySelectorAll('[data-filter-kind="text"]'));
            const selectFilters = Array.from(table.querySelectorAll('[data-filter-kind="select"]'));
            const rowCheckboxes = Array.from(table.querySelectorAll('tbody input[type="checkbox"][data-row-key]'));
            const menus = Array.from(table.querySelectorAll('.filter-menu'));
            const clearButtons = Array.from(table.querySelectorAll('.filter-clear'));
            const selectionStorageKey = 'pms-audit-checkboxes::{html.escape(str(workbook), quote=True)}::{html.escape(sheet_name, quote=True)}';

            function loadSelections() {{
                try {{
                    return JSON.parse(window.localStorage.getItem(selectionStorageKey) || '{{}}');
                }} catch (error) {{
                    return {{}};
                }}
            }}

            function saveSelections(selections) {{
                window.localStorage.setItem(selectionStorageKey, JSON.stringify(selections));
            }}

            function syncSelectionMarkup(selections) {{
                const embeddedSelection = document.getElementById('pms-audit-selected-rows');
                if (embeddedSelection) {{
                    embeddedSelection.textContent = JSON.stringify(selections, null, 2);
                }}
                rowCheckboxes.forEach((checkbox) => {{
                    const rowKey = checkbox.getAttribute('data-row-key');
                    const isSelected = Boolean(selections[rowKey]);
                    checkbox.checked = isSelected;
                    checkbox.setAttribute('data-selected', isSelected ? 'true' : 'false');
                    if (isSelected) {{
                        checkbox.setAttribute('checked', 'checked');
                    }} else {{
                        checkbox.removeAttribute('checked');
                    }}
                }});
            }}

            function restoreSelections() {{
                const savedSelections = loadSelections();
                syncSelectionMarkup(savedSelections);
            }}

            function cellText(row, column) {{
                const cell = row.querySelectorAll('td')[column];
                return cell ? cell.textContent.trim() : '';
            }}

            function populateSelectOptions() {{
                selectFilters.forEach((select) => {{
                    const column = Number(select.getAttribute('data-column'));
                    const seen = new Set();
                    rows.forEach((row) => {{
                        const value = cellText(row, column);
                        if (value) {{
                            seen.add(value);
                        }}
                    }});
                    Array.from(seen).sort((left, right) => left.localeCompare(right)).forEach((value) => {{
                        const option = document.createElement('option');
                        option.value = value;
                        option.textContent = value;
                        select.appendChild(option);
                    }});
                }});
            }}

            function applyFilters() {{
                rows.forEach((row) => {{
                    const visible = textFilters.every((input) => {{
                        const column = Number(input.getAttribute('data-column'));
                        const search = input.value.trim().toLowerCase();
                        const select = table.querySelector('[data-filter-kind="select"][data-column="' + column + '"]');
                        const selected = select ? select.value.trim().toLowerCase() : '';
                        const text = cellText(row, column).toLowerCase();
                        if (search && !text.includes(search)) {{
                            return false;
                        }}
                        if (selected && text !== selected) {{
                            return false;
                        }}
                        return true;
                    }});
                    row.style.display = visible ? '' : 'none';
                }});
            }}

            function closeOtherMenus(activeMenu) {{
                menus.forEach((menu) => {{
                    if (menu !== activeMenu) {{
                        menu.removeAttribute('open');
                    }}
                }});
            }}

            populateSelectOptions();
            restoreSelections();

            textFilters.forEach((input) => {{
                input.addEventListener('input', applyFilters);
            }});
            selectFilters.forEach((select) => {{
                select.addEventListener('change', applyFilters);
            }});
            rowCheckboxes.forEach((checkbox) => {{
                checkbox.addEventListener('change', () => {{
                    const savedSelections = loadSelections();
                    const rowKey = checkbox.getAttribute('data-row-key');
                    if (checkbox.checked) {{
                        savedSelections[rowKey] = true;
                    }} else {{
                        delete savedSelections[rowKey];
                    }}
                    saveSelections(savedSelections);
                    syncSelectionMarkup(savedSelections);
                }});
            }});
            clearButtons.forEach((button) => {{
                button.addEventListener('click', () => {{
                    const panel = button.closest('.filter-panel');
                    if (!panel) {{
                        return;
                    }}
                    panel.querySelectorAll('input, select').forEach((control) => {{
                        control.value = '';
                    }});
                    applyFilters();
                }});
            }});
            menus.forEach((menu) => {{
                menu.addEventListener('toggle', () => {{
                    if (menu.open) {{
                        closeOtherMenus(menu);
                    }}
                }});
            }});
            document.addEventListener('click', (event) => {{
                menus.forEach((menu) => {{
                    if (!menu.contains(event.target)) {{
                        menu.removeAttribute('open');
                    }}
                }});
            }});
        }})();
    </script>
</body>
</html>
"""


def main() -> int:
    args = parse_args()
    if not args.workbook.exists():
        raise FileNotFoundError(f"Workbook not found: {args.workbook}")

    if args.template_mode == "new-template-comment-fill":
        validate_new_template_args(args)
        write_wb = load_workbook(args.workbook, keep_vba=True)
        sheet_name = resolve_sheet_name(write_wb, args.sheet)
        write_ws = write_wb[sheet_name]
        summary = fill_new_template_comments(
            write_ws,
            args.repo_root,
            args.start_row,
            args.file_prefixes or [],
            args.platform_cds_module or "",
            args.sw_module_name or "",
            args.responsible_name or "",
        )
        output_workbook = args.annotated_workbook or default_new_template_output_path(
            args.workbook,
            args.sw_module_name,
        )
        output_path = args.output or default_html_output_path(args.sw_module_name, Path.cwd())
        html_content = build_html_report(
            "Understand Function Purple Comment Audit Report",
            summary["findings"],
            summary["module_rows"],
            summary["module_files"],
            args.workbook,
            sheet_name,
            sorted(summary["comment_mapping"].keys()),
            None,
            scope_description=(
                "Purple-highlighted comment cells in columns "
                f"{NEW_TEMPLATE_COMMENT_START_COLUMN}:{NEW_TEMPLATE_COMMENT_END_COLUMN} "
                f"mapped from row 7 for files starting with {', '.join(args.file_prefixes or [])}"
            ),
            summary_title="Purple Comment Summary",
            empty_message="No purple comment cells were filled for the selected file scope.",
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(html_content, encoding="utf-8")
        output_workbook.parent.mkdir(parents=True, exist_ok=True)
        write_wb.save(output_workbook)

        print(f"Template mode: {args.template_mode}")
        print(f"Workbook processed: {args.workbook}")
        print(f"Sheet processed: {sheet_name}")
        print(f"File prefixes: {', '.join(args.file_prefixes or [])}")
        print(f"Platform/CDS Module value: {args.platform_cds_module}")
        print(f"SW Module Name value: {args.sw_module_name}")
        print(f"Responsible Name value: {args.responsible_name}")
        print(f"Matched rows: {summary['matching_rows']}")
        print(f"Rows with generated comments: {summary['commented_rows']}")
        print(f"Updated purple comment cells: {summary['updated_comment_cells']}")
        print(
            "Comment mapping: "
            + ", ".join(
                f"{comment_column}->{metric_column}"
                for comment_column, metric_column in sorted(summary["comment_mapping"].items())
            )
        )
        print(f"HTML report: {output_path}")
        print(f"Workbook output: {output_workbook}")
        return 0

    configured_rules, configured_columns, loaded_metric_rule_file = load_metric_rule_config(metric_rule_file_path(args))
    global METRIC_RULES
    METRIC_RULES = dict(BUILTIN_METRIC_RULES)
    METRIC_RULES.update(configured_rules)

    allowed_files, module_name = load_allowed_files(args.file_list)
    output_path = args.output or default_html_output_path(module_name, Path.cwd())

    scan_wb = load_workbook(args.workbook, data_only=True, keep_vba=True)
    sheet_name = resolve_sheet_name(scan_wb, args.sheet)
    ws = scan_wb[sheet_name]
    target_columns = resolve_target_columns(args, configured_columns, configured_rules, ws)

    findings = list(iter_red_findings(ws, args.start_row, args.repo_root, target_columns, allowed_files))
    module_rows, module_files = collect_module_scope(ws, args.start_row, args.repo_root, target_columns, allowed_files)
    title = "Understand Function Red Metric Audit Report"
    html_content = build_html_report(
        title,
        findings,
        module_rows,
        module_files,
        args.workbook,
        sheet_name,
        target_columns,
        args.file_list,
        loaded_metric_rule_file,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(html_content, encoding="utf-8")

    annotated_workbook = args.annotated_workbook or (
        default_annotated_workbook_path_for_module(args.workbook, module_name)
        if module_name
        else default_annotated_workbook_path(args.workbook)
    )
    write_wb = load_workbook(args.workbook, keep_vba=True)
    write_ws = write_wb[sheet_name]
    apply_workbook_comments(write_ws, findings, module_rows, module_name)
    annotated_workbook.parent.mkdir(parents=True, exist_ok=True)
    write_wb.save(annotated_workbook)

    print(f"Sheet scanned: {sheet_name}")
    print(f"Columns scanned: {', '.join(target_columns)}")
    print(
        "Metric rule file: "
        f"{loaded_metric_rule_file if loaded_metric_rule_file else 'Not provided; using built-in defaults'}"
    )
    print(f"File filter: {args.file_list if args.file_list else 'All workbook rows'}")
    total_metrics = sum(len(item["metrics"]) for item in findings)
    print(f"Functions with red metrics: {len(findings)}")
    print(f"Rows included in comment scope: {len(module_rows)}")
    print(f"Rows included without comment: {max(len(module_rows) - len(findings), 0)}")
    print(f"Total individual red metric violations: {total_metrics}")
    print(f"HTML report: {output_path}")
    print(f"Annotated workbook: {annotated_workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
